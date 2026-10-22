"""导出报告共用的版式：页面、页眉页脚、标题区、摘要卡、表格、数据条、热力着色和证据引用。

人岗匹配报告和 AI 面试报告都从这里取颜色、字体和表格样式，各类报告外观保持一致。
这里只管版式：不读数据库，不改变报告里的分数、结论和证据。
"""
from __future__ import annotations

import io
import re
from bisect import bisect_right
from datetime import datetime
from itertools import accumulate

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips
from docx.table import _Cell

# 主色深绿，辅助浅色两档，警示色一种；其余为文字和边线的灰度。黑白打印时靠深浅和文字区分。
BRAND, LIGHT, PALE = "197465", "E7F5F2", "F2F9F8"
WARN, WARN_PALE = "C98D58", "FBF6F0"
INK, MUTED, BORDER = "1F2927", "5B6573", "C9DCD8"
BAR, TRACK = "5E9E93", "E4ECEA"
# 数据条的粗细（磅）和左端留白（twip）。
BAND_HEIGHT, BAND_INSET = 9, 110
FONT = "微软雅黑"

ORGANIZATION = "TalentFlow · 人力资源智能体"
FOOTER_NOTE = "AI辅助生成，须人工复核"
UNKNOWN_NAME = "待人工确认"

# A4，单位为 twip（1 厘米 = 567）：(页宽, 页高, 上, 下, 左右, 页眉距, 页脚距)。
PORTRAIT = (11906, 16838, 1361, 1304, 1134, 624, 510)
LANDSCAPE = (16838, 11906, 1304, 1247, 1020, 624, 510)

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
# 段落属性里排在边线（pBdr）之后的元素；Word 要求按这个顺序写。
_AFTER_BORDER = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                 "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                 "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                 "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                 "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def clean(value) -> str:
    """去掉 XML 不允许的控制字符，个别转写或解析文本不会导致导出失败。"""
    return _CONTROL.sub("", "" if value is None else str(value))


def person_label(name, source_filename="") -> str:
    """姓名未确认时显示“姓名待确认（简历文件名）”，与页面一致；文件名不带路径和扩展名。"""
    name = str(name or "").strip()
    if name and name != UNKNOWN_NAME:
        return name
    stem = re.split(r"[\\/]", str(source_filename or "").strip())[-1]
    stem = re.sub(r"\.[A-Za-z0-9]{1,5}$", "", stem).strip()
    return f"姓名待确认（{stem}）" if stem else "姓名待确认"


def now_text() -> str:
    return datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S")


# ---- 文字与段落 ------------------------------------------------------------------

def _set_fonts(properties):
    """中文、西文都用同一种字体；去掉模板里的主题字体，Word 和 WPS 才会按这里的设置显示。"""
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for key in list(fonts.attrib):
        del fonts.attrib[key]
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)


def style_run(run, size=10.5, bold=False, color=INK):
    _set_fonts(run._element.get_or_add_rPr())
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def _line(name, color, size=4, space=0):
    node = OxmlElement("w:" + name)
    values = (("val", "single"), ("sz", size), ("space", space), ("color", color)) if color else (("val", "nil"),)
    for key, value in values:
        node.set(qn("w:" + key), str(value))
    return node


def _shade(fill):
    node = OxmlElement("w:shd")
    for key, value in (("val", "clear"), ("color", "auto"), ("fill", fill)):
        node.set(qn("w:" + key), value)
    return node


def _write(p, text, size, bold, color, *, align=None, before=0, after=4, line=1.2, left=0, right=0,
           keep=False, borders=None, fill=None):
    """设置一个段落的边线、底色、间距并写入文字。borders 形如 {"left": (颜色, 粗细, 间距)}。"""
    properties = p._p.get_or_add_pPr()
    if borders:
        box = OxmlElement("w:pBdr")
        for name in ("top", "left", "bottom", "right"):
            if name in borders:
                box.append(_line(name, *borders[name]))
        properties.insert_element_before(box, *_AFTER_BORDER)
    if fill:
        properties.insert_element_before(_shade(fill), *_AFTER_BORDER[1:])
    layout = p.paragraph_format
    layout.space_before, layout.space_after, layout.line_spacing = Pt(before), Pt(after), line
    if left:
        layout.left_indent = Twips(left)
    if right:
        layout.right_indent = Twips(right)
    if keep:
        layout.keep_with_next = True
    if align:
        p.alignment = _ALIGN[align]
    for part in ([(text, size, bold, color)] if isinstance(text, str) else text):
        style_run(p.add_run(clean(part[0])), *part[1:])
    return p


def paragraph(target, text="", *, size=10.5, bold=False, color=INK, style=None, **layout):
    """在文档或单元格里加一段。text 为字符串，或 [(文字, 字号, 加粗, 颜色), ...] 表示同一段里的几种字样。"""
    return _write(target.add_paragraph(style=style), text, size, bold, color, **layout)


def heading(doc, text, level=1):
    """一级标题带浅色底和左侧色条，二级标题带下边线，三级标题只加粗；都与后面的内容留在同一页。"""
    if level == 1:
        return paragraph(doc, text, size=13, bold=True, color=BRAND, style="Heading 1", before=14, after=7, line=1.15,
                         left=200, keep=True, borders={"left": (BRAND, 24, 7)}, fill=LIGHT)
    if level == 2:
        return paragraph(doc, text, size=11.5, bold=True, color=INK, style="Heading 2", before=10, after=5, line=1.15,
                         keep=True, borders={"bottom": (BORDER, 6, 3)})
    return paragraph(doc, text, size=10.5, bold=True, color=BRAND, style="Heading 3", before=8, after=3, line=1.15,
                     keep=True)


def bullets(doc, values, size=10.5):
    for value in values:
        paragraph(doc, str(value), size=size, style="List Bullet", after=3, line=1.25)


def numbered(doc, values, size=10.5):
    """带序号的条目：用 Word 的自动编号，每个列表都从 1 开始，换行后的文字与首行对齐。"""
    numbering = doc.part.numbering_part.element
    base = numbering.num_having_numId(doc.styles["List Number"].element.pPr.numPr.numId.val)
    number = numbering.add_num(base.abstractNumId.val)
    number.add_lvlOverride(ilvl=0).add_startOverride(1)
    for value in values:
        p = paragraph(doc, str(value), size=size, style="List Number", after=3, line=1.25)
        properties = p._p.get_or_add_pPr().get_or_add_numPr()
        properties.get_or_add_ilvl().val = 0
        properties.get_or_add_numId().val = number.numId


def note(target, text, *, label="", color=WARN, fill=WARN_PALE, size=9.5, after=6):
    """提示框：左侧色条加浅底，用于免责声明和计算口径说明。"""
    hidden = (fill, 4, 4)
    runs = ([(label, size, True, INK)] if label else []) + [(text, size, False, MUTED)]
    return paragraph(target, runs, after=after, line=1.3, left=215, right=90,
                     borders={"top": hidden, "left": (color, 18, 8), "bottom": hidden, "right": hidden}, fill=fill)


def quote(target, text, source, *, size=9.5, after=6):
    """证据引用：左边线加缩进，上面是原文，下面是出处（文件 · 位置）。"""
    line = {"left": (BRAND, 12, 9)}
    paragraph(target, text, size=size, after=1, line=1.3, left=420, keep=True, borders=line)
    return paragraph(target, "出处：" + source, size=8.5, color=MUTED, after=after, line=1.2, left=420, borders=line)


# ---- 页面、页眉页脚、标题区 --------------------------------------------------------

def content_width(doc) -> int:
    """版心宽度（twip）。"""
    section = doc.sections[0]
    return round((section.page_width - section.left_margin - section.right_margin) / 635)


def _field(p, code, size, color):
    """页码用 Word 域，打开或打印时由 Word 计算。"""
    for kind in ("begin", "code", "separate", "result", "end"):
        run = style_run(p.add_run("1" if kind == "result" else ""), size, False, color)
        if kind == "code":
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = f" {code} "
        elif kind != "result":
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        else:
            continue
        run._element.append(node)


def _header_footer(section, report_type, width):
    header = section.header.paragraphs[0]
    header.style = "Normal"
    _write(header, [(ORGANIZATION, 9, True, BRAND), ("\t" + report_type, 9, False, MUTED)], 9, False, MUTED,
           after=0, line=1.0, borders={"bottom": (BRAND, 6, 5)})
    header.paragraph_format.tab_stops.add_tab_stop(Twips(width), WD_TAB_ALIGNMENT.RIGHT)
    page = section.footer.paragraphs[0]
    page.style = "Normal"
    _write(page, "第 ", 9, False, MUTED, align="center", after=0, line=1.0, borders={"top": (BORDER, 4, 5)})
    _field(page, "PAGE", 9, MUTED)
    style_run(page.add_run(" 页 / 共 "), 9, False, MUTED)
    _field(page, "NUMPAGES", 9, MUTED)
    style_run(page.add_run(" 页"), 9, False, MUTED)
    paragraph(section.footer, FOOTER_NOTE, size=8, color=MUTED, style="Normal", align="center", after=0, line=1.0)


def _prepare_styles(doc):
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None and defaults.find(".//" + qn("w:rPr")) is not None:
        _set_fonts(defaults.find(".//" + qn("w:rPr")))
        for language in defaults.iter(qn("w:lang")):
            language.set(qn("w:eastAsia"), "zh-CN")
    # 按中文文档排版：行末标点压缩，避免为了容纳标点而挤压同一行里的制表位。
    for tag, key, value in (("w:characterSpacingControl", "w:val", "compressPunctuation"),
                            ("w:themeFontLang", "w:eastAsia", "zh-CN")):
        for node in doc.settings.element.iter(qn(tag)):
            node.set(qn(key), value)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number", "Header", "Footer"):
        style = doc.styles[name]
        _set_fonts(style.element.get_or_add_rPr())
        style.font.color.rgb = RGBColor.from_string(INK)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)


def new_document(title, report_type="", *, landscape=False):
    """新建报告：A4 页面、页眉（左侧单位名称，右侧报告类型）、页脚（页码和复核提示）和报告标题。"""
    doc = Document()
    section = doc.sections[0]
    width, height, top, bottom, side, header, footer = LANDSCAPE if landscape else PORTRAIT
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Twips(width), Twips(height)
    section.top_margin, section.bottom_margin = Twips(top), Twips(bottom)
    section.left_margin = section.right_margin = Twips(side)
    section.header_distance, section.footer_distance = Twips(header), Twips(footer)
    _prepare_styles(doc)
    _header_footer(section, clean(report_type or title), width - 2 * side)
    properties = doc.core_properties
    properties.title, properties.author, properties.comments = clean(title), ORGANIZATION, ""
    properties.created = properties.modified = datetime.now()
    paragraph(doc, title, size=20, bold=True, color=INK, style="Title", after=3, line=1.15)
    return doc


def info_line(doc, items, columns=2):
    """标题下的基本信息：对象、岗位、生成时间等。每行两项，用制表位对齐，各项不会从中间断开；

    下方一条主色细线与正文分开。
    """
    runs = []
    for index, (label, value) in enumerate(items):
        if index:
            runs.append(("\n" if index % columns == 0 else "\t", 9.5, False, MUTED))
        runs += [(f"{label}：", 9.5, False, MUTED), (str(value), 9.5, False, INK)]
    p = paragraph(doc, runs, after=10, line=1.35, borders={"bottom": (BRAND, 12, 6)})
    for column in range(1, columns):
        p.paragraph_format.tab_stops.add_tab_stop(Twips(content_width(doc) * column // columns))
    return p


def save_bytes(doc) -> bytes:
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


# ---- 表格 ------------------------------------------------------------------------

def spread(shares, total) -> list[int]:
    """按比例把总宽分给各列，取整后的差额并入最后一列。"""
    whole = float(sum(shares))
    widths = [int(total * share / whole) for share in shares]
    widths[-1] += total - sum(widths)
    return widths


def cell(text, width, *, size=9.5, bold=False, color=INK, fill=None, align="left", edges=None, band=None, inset=0):
    """一格的内容和样式。text 为字符串，或段落列表（每段是 [(文字, 字号, 加粗, 颜色), ...]）。

    edges 指定个别边线，如 {"left": (颜色, 粗细)}；值为 None 表示不画这条线。
    band 是数据条的颜色：这一格不写字，只画一条横向色带，inset 是色带左端留出的空白。
    """
    return {"text": text, "width": int(width), "size": size, "bold": bold, "color": color,
            "fill": fill, "align": align, "edges": edges or {}, "band": band, "inset": inset}


def _cell_padding(top, left, bottom, right, tag="w:tcMar"):
    padding = OxmlElement(tag)
    for name, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        item = OxmlElement("w:" + name)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")
        padding.append(item)
    return padding


def _table_frame(table, edges, border, margins, full):
    properties = table._tbl.tblPr
    for child in list(properties):
        properties.remove(child)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(edges[-1]))
    width.set(qn("w:type"), "dxa")
    # 宽度等于版心的表格居中放置，两侧与正文对齐；较窄的表格靠左，左边线与其他表格对齐。
    position = OxmlElement("w:jc" if full else "w:tblInd")
    if full:
        position.set(qn("w:val"), "center")
    else:
        position.set(qn("w:w"), str(margins[1]))
        position.set(qn("w:type"), "dxa")
    box = OxmlElement("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        box.append(_line(name, border))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    padding = _cell_padding(margins[0], margins[1], margins[0], margins[1], "w:tblCellMar")
    for node in (width, position, box, layout, padding):
        properties.append(node)
    for column, left, right in zip(table._tbl.tblGrid.gridCol_lst, [0, *edges], edges):
        column.set(qn("w:w"), str(right - left))


def _table_cell(table, item, span, keep):
    tc = OxmlElement("w:tc")
    properties = OxmlElement("w:tcPr")
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(item["width"]))
    width.set(qn("w:type"), "dxa")
    properties.append(width)
    if span > 1:
        merged = OxmlElement("w:gridSpan")
        merged.set(qn("w:val"), str(span))
        properties.append(merged)
    if item["edges"]:
        box = OxmlElement("w:tcBorders")
        for name in ("top", "left", "bottom", "right"):
            if name in item["edges"]:
                box.append(_line(name, *(item["edges"][name] or (None,))))
        properties.append(box)
    if item["fill"]:
        properties.append(_shade(item["fill"]))
    if item["band"]:
        properties.append(_cell_padding(0, item["inset"], 0, 0))
    middle = OxmlElement("w:vAlign")
    middle.set(qn("w:val"), "center")
    properties.append(middle)
    tc.append(properties)
    tc.append(OxmlElement("w:p"))
    target = _Cell(tc, table)
    if item["band"]:
        # 色带是一个带底色的空段落，高度固定，在格内垂直居中，行再高也不会变粗。
        _write(target.paragraphs[0], "", 1, False, INK, after=0, line=Pt(BAND_HEIGHT), fill=item["band"], keep=keep)
        return tc
    parts = [item["text"]] if isinstance(item["text"], str) else item["text"]
    for index, part in enumerate(parts):
        p = target.paragraphs[0] if index == 0 else target.add_paragraph()
        _write(p, part, item["size"], item["bold"], item["color"], align=item["align"], after=0, line=1.15, keep=keep)
    return tc


def grid_table(doc, rows, *, header_rows=0, border=BORDER, margins=(75, 115), split=False, together=None):
    """按每格宽度排布的表格。各行的分格位置可以不同（数据条、摘要卡用到），网格列取全部分格位置的并集。

    固定列宽，长文本在格内换行；表头行跨页重复。
    分页：一行默认不拆到两页；小表格（纵向页面不超过 5 行，横向页面不超过 3 行，或 together=True）
    整体留在同一页；其余表格可以在行与行之间分页，但开头（表头加两行）和最后两行各自不分开。
    单元格里是整段原文的表格传 split=True，允许在页末拆行，避免留下大块空白。
    """
    edges = sorted({edge for row in rows for edge in accumulate(item["width"] for item in row)})
    table = doc.add_table(rows=0, cols=len(edges))
    _table_frame(table, edges, border, margins, edges[-1] >= content_width(doc))
    if together is None:
        section = doc.sections[0]
        together = not split and len(rows) <= (5 if section.page_height > section.page_width else 3)
    lead = header_rows if split else max(header_rows + 1, 2)
    for index, row in enumerate(rows):
        tr = OxmlElement("w:tr")
        properties = OxmlElement("w:trPr")
        if not split or index < header_rows:
            properties.append(OxmlElement("w:cantSplit"))
        if index < header_rows:
            properties.append(OxmlElement("w:tblHeader"))
        tr.append(properties)
        keep = index < len(rows) - 1 and (together or index < lead or (not split and index == len(rows) - 2))
        start = 0
        for item in row:
            end = start + item["width"]
            tr.append(_table_cell(table, item, bisect_right(edges, end) - bisect_right(edges, start), keep))
            start = end
        table._tbl.append(tr)
    # 表格后留一小段空白；两个表格相邻时也不会被 Word 合并成一个。
    paragraph(doc, "", size=2, after=3, line=1.0)
    return table


def header_cells(headers, widths, centered=(), size=9.5):
    return [cell(text, width, size=size, bold=True, color=INK, fill=LIGHT, align="center" if index in centered else "left")
            for index, (text, width) in enumerate(zip(headers, widths))]


def table(doc, headers, records, widths, centered=(), *, size=9.5, split=False):
    """统一样式的表格：表头浅色底，细边线，列宽按比例铺满版心。"""
    widths, centered = spread(widths, content_width(doc)), set(centered)
    rows = [header_cells(headers, widths, centered, size)]
    for record in records:
        rows.append([cell(str(text), width, size=size, align="center" if index in centered else "left")
                     for index, (text, width) in enumerate(zip(record, widths))])
    return grid_table(doc, rows, header_rows=1, split=split)


def pairs(doc, rows, *, label=1600, split=False, together=None):
    """标签与内容对照的表格：每行一到三组，标签格带浅色底。rows 的每一行是 [(标签, 内容), ...]。"""
    total, lines = content_width(doc), []
    for row in rows:
        # 一行三组时标签都是短词，标签格收窄，把宽度留给内容。
        name_width = label if len(row) < 3 else 1250
        room = (total - name_width * len(row)) // len(row)
        line = [item for name, value in row
                for item in (cell(name, name_width, bold=True, fill=LIGHT), cell(str(value), room))]
        line[-1]["width"] += total - sum(item["width"] for item in line)
        lines.append(line)
    return grid_table(doc, lines, split=split, together=together)


def bar_cells(width, score, text, *, maximum=100, number=900, size=9.5):
    """数据条：把一格拆成着色、留白和数字三格，着色宽度按得分比例（取 2% 的整数倍）。

    没有得分（score 为 None）时整格写说明文字，不画条，也不按 0 分处理。
    列太窄画不下数据条时只写数字。
    """
    if score is None:
        return [cell(text, width, size=size, color=MUTED, align="center")]
    room = width - number - BAND_INSET
    if room < 400:
        return [cell(text, width, size=size, bold=True, align="center")]
    filled = room * round(min(max(float(score) / maximum, 0), 1) * 50) // 50
    parts = [cell("", part, band=color) for part, color in ((filled, BAR), (room - filled, TRACK)) if part]
    parts[0].update(width=parts[0]["width"] + BAND_INSET, inset=BAND_INSET)
    parts.append(cell(text, number, size=size, bold=True, align="center"))
    # 条与数字属于同一列，中间不画竖线。
    for left, right in zip(parts, parts[1:]):
        left["edges"]["right"] = right["edges"]["left"] = None
    return parts


def heat_fill(score, maximum=100):
    """热力底色：得分越高颜色越深，只用浅色，黑色文字保持清晰；没有得分时不着色。"""
    if score is None:
        return None
    depth = 0.06 + 0.36 * min(max(float(score) / maximum, 0), 1)
    return "".join(f"{round(255 - (255 - channel) * depth):02X}" for channel in (0x19, 0x74, 0x65))


def metric(label, value, unit="", size=24):
    """摘要卡里的大号数字：标签、数值和单位写在同一段。"""
    return [[(f"{label}：", 10, False, MUTED), (str(value), size, True, BRAND), (unit, 10, False, MUTED)]]


def fact(label, value, size=11):
    """摘要卡里的一项：上面是小字标签，下面是加粗的内容。"""
    return [[(label, 9, False, MUTED)], [(str(value), size, True, INK)]]


def inline_fact(label, value, size=13):
    """摘要卡里标签和内容写在同一段的一项。"""
    return [[(f"{label}：", 10, False, MUTED), (str(value), size, True, INK)]]


def card(doc, rows):
    """结果摘要卡：带底色的表格，左侧一条主色粗线。rows 的每一行是 [(宽度占比, 段落列表), ...]。"""
    total, lines = content_width(doc), []
    for number, row in enumerate(rows):
        widths = spread([share for share, _ in row], total)
        line = [cell(parts, width, fill=LIGHT if number == 0 else PALE) for (_, parts), width in zip(row, widths)]
        line[0]["edges"]["left"] = (BRAND, 36)
        lines.append(line)
    return grid_table(doc, lines, border="B5D3CD", margins=(120, 190), together=True)
