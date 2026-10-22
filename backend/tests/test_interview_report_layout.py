"""AI 面试报告、综合报告和候选人对比的版式：页眉页脚、摘要卡、数据条、雷达图和热力着色。"""
import io

import pytest
from docx import Document
from docx.oxml.ns import qn

from backend.app.interview import docx_export, mock_engine
from backend.app.interview.docx_export import comparison_docx, radar_png, report_docx
from matching_core.services import report_style as style

METADATA = {"candidate_name": "版式测试人员", "candidate_code": "LAYOUT-001", "job_name": "人事专员",
            "batch_name": "版式测试批次", "version": 2}


def content(scores=(4, 3, 5, 0, 2, 4), missing=()):
    """六题均已作答的虚构报告；missing 中的题目回答过短，没有形成评分。"""
    answers = [{"sequence": index, "answer_code": f"A{index}", "question": f"第{index}题",
                "transcript": "没有例子。" if index in missing else transcript}
               for index, transcript in enumerate(mock_engine.DEFAULT_TRANSCRIPTS, start=1)]
    return mock_engine.build_report("版式测试人员", answers, list(scores))


def text_of(doc):
    return "\n".join([*(p.text for p in doc.paragraphs),
                      *(cell.text for table in doc.tables for row in table.rows for cell in row.cells)])


def table_headed(doc, first_header):
    return next(table for table in doc.tables if table.rows[0].cells[0].text == first_header)


def bands(row):
    return [shade.get(qn("w:fill")) for shade in row._tr.iter(qn("w:shd")) if shade.getparent().tag == qn("w:pPr")]


def cell_fill(cell):
    shade = cell._tc.tcPr.find(qn("w:shd"))
    return None if shade is None else shade.get(qn("w:fill"))


def item(code, name, scores, match, stage_two=None):
    dimensions = [entry[0] for entry in mock_engine.QUESTIONS]
    return {"candidate": {"id": 1, "code": code, "name": name}, "status": "confirmed" if scores else "not_generated",
            "match_score": match, "scores": dict(zip(dimensions, scores)) if scores else {}, "risks": ["转写需核对"] if scores else [],
            "hard_indicators": {"available": False, "reason": "候选人尚未关联人员库档案", "facts": [], "hard_checks": []},
            "stage_two": stage_two}


def test_interview_report_has_page_furniture_summary_card_and_disclaimer():
    doc = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(), METADATA)))
    section = doc.sections[0]
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    header = section.header.paragraphs[0].text
    assert "TalentFlow · 人力资源智能体" in header and "TalentFlowAI面试报告" in header
    assert [node.text.strip() for node in section.footer._element.iter(qn("w:instrText"))] == ["PAGE", "NUMPAGES"]
    assert "AI辅助生成，须人工复核" in "\n".join(p.text for p in section.footer.paragraphs)
    assert doc.paragraphs[0].style.name == "Title" and doc.paragraphs[0].text == "TalentFlowAI面试报告"
    basics = doc.paragraphs[1].text
    for expected in ("候选人：版式测试人员", "候选人编号：LAYOUT-001", "应聘岗位：人事专员", "招聘批次：版式测试批次", "导出时间："):
        assert expected in basics
    card = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    for expected in ("辅助匹配度：", "已形成评分 6/6 项", "v2"):
        assert expected in card
    # 报告只在人工确认后才能导出，不再写报告状态。
    assert "报告状态" not in card and "已人工确认" not in text_of(doc)
    assert {cell_fill(cell) for row in doc.tables[0].rows for cell in row.cells} <= {style.LIGHT, style.PALE}
    assert "AI辅助生成，须人工复核。本文档用于辅助人工评估，不构成录用或淘汰决定。" in text_of(doc)


def test_scored_dimensions_have_bars_and_unscored_ones_say_evidence_is_insufficient():
    doc = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(missing=(2,)), METADATA)))
    table = table_headed(doc, "评价维度")
    rows = {row.cells[0].text: row for row in table.rows[1:]}
    assert bands(rows["岗位专业能力"]) == [style.BAR, style.TRACK]
    assert "4 / 5" in [cell.text for cell in rows["岗位专业能力"].cells]
    # 满分整条着色；0 分是有效评分，画空条并保留数字。
    assert bands(rows["沟通表达"]) == [style.BAR] and bands(rows["抗压能力"]) == [style.TRACK]
    assert "0 / 5" in [cell.text for cell in rows["抗压能力"].cells]
    # 没有评分的维度不画条，也不写成 0 分。
    assert bands(rows["逻辑思维"]) == []
    texts = [cell.text for cell in rows["逻辑思维"].cells]
    assert "证据不足" in texts and not any(text.startswith("0") for text in texts)
    assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


def test_report_without_missing_evidence_never_mentions_it():
    data = report_docx("TalentFlowAI面试报告", content(), METADATA)
    doc = Document(io.BytesIO(data))
    assert "证据不足" not in text_of(doc)
    assert "0级" in text_of(doc) and "0 / 5" in text_of(doc)


def test_radar_chart_is_added_only_when_enough_dimensions_are_scored():
    full = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(), METADATA)))
    assert len(full.inline_shapes) == 1
    # 图片宽度不超过版心。
    section = full.sections[0]
    assert full.inline_shapes[0].width <= section.page_width - section.left_margin - section.right_margin
    assert "六维辅助评分雷达图" in text_of(full)
    few = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(missing=(1, 2, 3, 4)), METADATA)))
    assert len(few.inline_shapes) == 0 and "雷达图" not in text_of(few)
    standalone = Document(io.BytesIO(report_docx("TalentFlow人工面试AI旁听综合报告", {
        "mode": "standalone", "stage_two_summary": "人工面试共留存 1 段有效转写。",
        "stage_two_evidence": [{"id": "c-1", "speaker": "candidate", "text": "我负责过三年的招聘工作。"}]}, METADATA)))
    assert len(standalone.inline_shapes) == 0


def test_radar_picture_is_png_and_skips_unusable_input(monkeypatch):
    labels = [entry[0] for entry in mock_engine.QUESTIONS]
    assert radar_png(labels, [4, 3, 5, 0, 2, 4])[:8] == b"\x89PNG\r\n\x1a\n"
    assert radar_png(labels, [4, None, 5, 0, None, 4])[:8] == b"\x89PNG\r\n\x1a\n"
    assert radar_png(labels, [4, None, None, None, None, 4]) is None
    assert radar_png(labels[:2], [4, 3]) is None

    def broken(*args, **kwargs):
        raise RuntimeError("绘图失败")

    # 绘图出错时不放图，报告照常导出。
    monkeypatch.setattr(docx_export, "_draw_radar", broken)
    assert radar_png(labels, [4, 3, 5, 0, 2, 4]) is None
    doc = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(), METADATA)))
    assert len(doc.inline_shapes) == 0 and "六维评价" in text_of(doc)


@pytest.mark.parametrize("name, shown", [("待人工确认", "姓名待确认"), ("版式测试人员", "版式测试人员")])
def test_unconfirmed_candidate_name_is_not_shown_as_a_name(name, shown):
    doc = Document(io.BytesIO(report_docx("TalentFlowAI面试报告", content(), dict(METADATA, candidate_name=name))))
    assert f"候选人：{shown}\t" in doc.paragraphs[1].text
    assert "候选人：待人工确认" not in doc.paragraphs[1].text


def test_comparison_is_landscape_with_shaded_scores_and_match_bars():
    items = [item("CMP-001", "对比人员甲", (5, 4, 3, 2, 1, 0), 80.0, {"status": "confirmed"}),
             item("CMP-002", "待人工确认", (3, None, 3, 3, 3, 3), 60.0),
             item("CMP-003", "对比人员丙", None, None, {"status": "confirmed"})]
    doc = Document(io.BytesIO(comparison_docx("人事专员 · 版式测试批次", items)))
    section = doc.sections[0]
    assert section.page_width > section.page_height and round(section.page_width.cm, 1) == 29.7
    assert "候选人对比报告" in section.header.paragraphs[0].text
    assert [node.text.strip() for node in section.footer._element.iter(qn("w:instrText"))] == ["PAGE", "NUMPAGES"]
    text = text_of(doc)
    assert "对比结果仅供人工初筛参考，不构成自动录用或淘汰决定。" in text
    assert "姓名待确认" in text and "待人工确认（" not in text and "None" not in text
    card = "\n".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "对比人数：3 人" in card and "已有六维评分：2 人" in card and "80%" in card
    scores = table_headed(doc, "候选人")
    first, second, third = scores.rows[1], scores.rows[2], scores.rows[3]
    fills = [cell_fill(cell) for cell in first.cells[1:7]]
    assert fills == [style.heat_fill(value, 5) for value in (5, 4, 3, 2, 1, 0)]
    assert sorted(fills) == fills and len(set(fills)) == 6
    # 没有评分的格子不着色；岗位匹配度画数据条，没有匹配度的不画。
    assert second.cells[2].text == "证据不足" and cell_fill(second.cells[2]) is None
    assert bands(first) == [style.BAR, style.TRACK] and "80%" in [cell.text for cell in first.cells]
    assert bands(third) == [] and third.cells[-1].text == "—"
    # 表格宽度等于版心，不超出页面。
    room = round((section.page_width - section.left_margin - section.right_margin) / 635)
    assert {int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) for table in doc.tables} == {room}


def test_comparison_without_first_stage_scores_says_so():
    doc = Document(io.BytesIO(comparison_docx("人事专员", [item("CMP-009", "仅旁听人员", None, None, {"status": "confirmed"})])))
    text = text_of(doc)
    assert "没有六维评分" in text and "None%" not in text and "硬指标" in text
    # 摘要卡之后只有旁听情况、风险和硬指标三张表，没有评分对比表；不写报告状态。
    assert [table.rows[0].cells[-1].text for table in doc.tables[1:]] == ["人工面试旁听", "风险提示", "岗位硬性条件核验"]
    assert "已做，有综合报告" in text and "报告状态" not in text and "已人工确认" not in text
