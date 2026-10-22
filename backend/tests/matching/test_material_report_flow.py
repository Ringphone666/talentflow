"""Material corrections must reach fresh reports in both matching directions."""
from pathlib import Path
from . import helpers

import pytest
from docx import Document

from .fixture_data import init_db
from matching_core.services import batch_reports, batches, workflow


def _document_text(path):
    document = Document(path)
    return '\n'.join([*(p.text for p in document.paragraphs),
                      *(c.text for t in document.tables for r in t.rows for c in r.cells)])


@pytest.mark.parametrize('fit', [False, True])
def test_edit_reanalysis_reapproval_and_export_use_current_material(tmp_path, monkeypatch, fit):
    database = tmp_path / 'material-flow.db'
    init_db(database)
    monkeypatch.setattr(batch_reports, 'REPORT_DIR', tmp_path / 'exports')

    def local_assessment(score, parsed, job, *args):
        score.update(scoring_method='ai_semantic_v1',
                     dimensions=[{'dimension': name, 'weight': weight, 'score': 80,
                                  'reason': '虚构材料回归验证', 'evidence': []}
                                 for name, weight in job['weights'].items()],
                     hard_checks=[], reasoning_analysis={'status': '调用成功', 'model': 'test'})
        score.pop('requires_reassessment', None)
        return score

    monkeypatch.setattr(workflow, 'analyze_score', local_assessment)
    source = tmp_path / '虚构测试简历.txt'
    source.write_text('姓名：原姓名\n候选人编号：EDIT-001\n教育经历：本科\n工作经历：开发项目三年\n'
                      '项目经历：企业管理系统\n专业技能：Python\n任职时间：三年\n面试评价：表达清楚', encoding='utf-8')
    original_bytes = source.read_bytes()
    job_ids = [job['id'] for job in workflow.get_jobs(database)[:2]]
    task_id = (workflow.create_fit_task(job_ids, [source.name], '招聘人员', database) if fit else
               workflow.create_task(job_ids[0], [source.name], '招聘人员', database))
    if fit:
        candidate_id = workflow.process_fit_task(task_id, [source], '招聘人员', db_path=database)[0]
    else:
        candidate_id = workflow.process_task(task_id, job_ids[0], [source], '招聘人员', db_path=database)[0]
    candidate = workflow.get_candidate(candidate_id, database)
    helpers.update_candidate_fields(candidate_id, {key: field['value'] for key, field in candidate['parsed']['fields'].items()},
                                     '招聘人员', database)
    if fit:
        workflow.save_fit_decision(candidate_id, job_ids[0], '原审核意见', '推荐至该岗位', '招聘人员', database)
    else:
        workflow.save_review(candidate_id, '原审核意见', '进入下一轮', '招聘人员', database)
    assert batches.batch_snapshot(task_id, database)['rows'][0]['审核状态'] == '通过'

    workflow.update_candidate_material(candidate_id, {'education': '人工核实为硕士',
        'interview_evaluation': '已核对项目职责并补充面试纪要', 'additional_material': '人工补充的项目材料'},
        '招聘人员', database, full_name='修正姓名', confirmed=False)
    snapshot = batches.batch_snapshot(task_id, database)
    current = snapshot['rows'][0]
    assert current['姓名'] == '修正姓名'
    assert current['分析状态'] == '资料已更新，待重分析'
    assert current['审核状态'] == '需重新审核' and current['资料确认'] == '待确认'
    assert current['审核人'] == current['审核时间'] == ''
    with pytest.raises(ValueError):
        batch_reports.export_single_report(task_id, candidate_id, '招聘人员', database)
    assert workflow.reanalyze_task(task_id, '招聘人员', database) == [candidate_id]
    assert batches.batch_snapshot(task_id, database)['rows'][0]['审核状态'] == '待审核'
    workflow.update_candidate_material(candidate_id, {}, '招聘人员', database, confirmed=True)
    with pytest.raises(ValueError):
        batch_reports.export_single_report(task_id, candidate_id, '招聘人员', database)
    if fit:
        workflow.save_fit_decision(candidate_id, job_ids[0], '按修正资料重新审核', '推荐至该岗位', '招聘人员', database)
    else:
        workflow.save_review(candidate_id, '按修正资料重新审核', '进入下一轮', '招聘人员', database)
    path = batch_reports.export_single_report(task_id, candidate_id, '招聘人员', database)
    text = _document_text(path)
    for value in ('修正姓名', '人工核实为硕士', '已核对项目职责并补充面试纪要', '人工补充的项目材料'):
        assert value in text
    assert '候选人姓名：原姓名' not in text
    assert Path(workflow.get_candidate(candidate_id, database)['original_path']).read_bytes() == original_bytes
