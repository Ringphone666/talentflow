import pytest
from docx import Document
from matching_core.services import job_extraction as ai
from matching_core.services.model_adapter import ModelError
from matching_core.services.parser import extract_minutes


@pytest.fixture
def model(monkeypatch):
    monkeypatch.setattr(ai, 'model_settings', lambda: dict(provider='external', confirmed=True, key='test-secret', url='https://example.invalid', http_allowed=False))
    result = dict(metadata=dict(name='销售专员', sequence='营销类', level='P3', jd_version=''),
                  responsibilities=['负责客户开发'], requirements=['具备沟通能力'], hard_conditions=[], dimensions=[])
    def analyze(self, text, job):
        assert self.model == 'chat-primary' and self.task in ('job_evidence', 'job_hard_conditions')
        assert '销售专员' in text
        if self.task == 'job_hard_conditions':
            return {'hard_conditions': result['hard_conditions']}
        return result
    monkeypatch.setattr(ai.ExternalModelAdapter, 'analyze', analyze)
    return result


def test_word_tables_and_ai_extraction(tmp_path, model):
    doc = Document()
    doc.add_paragraph('岗位说明书')
    table = doc.add_table(rows=5, cols=2)
    for row, values in zip(table.rows, [('岗位名称','销售专员'),('岗位类别','营销类'),('职级','P3'),('岗位职责','负责客户开发'),('任职要求','具备沟通能力')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / 'job.docx'
    doc.save(path)
    assert '负责客户开发' in extract_minutes(path).text
    parsed = ai.extract_job_with_ai(path)
    assert parsed['metadata']['name'] == '销售专员'
    assert parsed['extraction_model'] == 'chat-primary'


@pytest.mark.parametrize('failure', ['invented', 'network', 'multiple'])
def test_invalid_or_failed_extraction_is_not_silently_replaced(tmp_path, monkeypatch, model, failure):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3 负责客户开发 具备沟通能力', encoding='utf-8')
    if failure == 'invented':
        model['requirements'] = ['必须具备十年经验']
    elif failure == 'multiple':
        model.clear()
        model['multiple_jobs'] = True
    else:
        def fail(*args):
            raise ModelError('连接失败')
        monkeypatch.setattr(ai.ExternalModelAdapter, 'analyze', fail)
    if failure == 'invented':
        result = ai.extract_job_with_ai(path)
        assert result['metadata']['name'] == '销售专员'
        assert result['requirements'] == []
        assert '任职要求' in result['warning']
    else:
        with pytest.raises(ValueError, match='chat-primary'):
            ai.extract_job_with_ai(path)


def test_partial_invalid_fields_keep_verified_evidence_and_original_punctuation(tmp_path, model):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员（区域）\n营销类\nP3\n负责客户开发\n具备沟通能力', encoding='utf-8')
    model['metadata']['name'] = '销售专员(区域)'
    model['metadata']['level'] = 'P8'
    model['responsibilities'] = [3, 999]
    model['requirements'] = [4]
    result = ai.extract_job_with_ai(path)
    assert result['metadata']['name'] == '销售专员（区域）'
    assert result['metadata']['level'] == ''
    assert result['responsibilities'] == ['负责客户开发']
    assert result['requirements'] == ['具备沟通能力']
    assert '职级' in result['warning'] and '岗位职责' in result['warning']


def test_extracts_source_scorecard_weights_with_original_evidence(tmp_path, model):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3\n负责客户开发\n具备沟通能力\n评分标准\n客户开发 | 70% | 客户拜访\n沟通能力 | 30% | 商务沟通', encoding='utf-8')
    model['dimensions'] = [dict(name='客户开发', weight=70, keywords=['客户拜访'], evidence_ids=[4]), dict(name='沟通能力', weight=30, keywords=['商务沟通'], evidence_ids=[5])]
    result = ai.extract_job_with_ai(path)
    assert result['weights'] == {'客户开发': 70, '沟通能力': 30}
    assert result['dimensions'][0] == {'name': '客户开发', 'keywords': ['客户拜访']}
    model['dimensions'][0]['weight'] = 99
    assert '客户开发' not in ai.extract_job_with_ai(path)['weights']


def test_hard_conditions_keep_only_cited_clauses_and_original_logic(tmp_path, model):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3\n负责客户开发\n具备沟通能力\n基本学历 | 本科 | 所学专业 | 营销类优先\n工作经验 | 3年以上销售经验或持有职业资格证', encoding='utf-8')
    model['hard_conditions'] = [
        {'text': '基本学历 | 本科', 'evidence_ids': [3]},
        {'text': '工作经验 | 3年以上销售经验或持有职业资格证', 'evidence_ids': [4]},
        {'text': '基本学历 | 本科', 'evidence_ids': [3]},
    ]
    result = ai.extract_job_with_ai(path)
    assert result['hard_conditions'] == ['基本学历 | 本科', '工作经验 | 3年以上销售经验或持有职业资格证']
    assert not result['warning']


@pytest.mark.parametrize('entry', [
    {'text': '必须博士', 'evidence_ids': [1]},
    {'text': '基本学历 | 本科', 'evidence_ids': [999]},
    {'text': '基本学历 | 本科', 'evidence_ids': [0]},
    {'text': '基本学历 | 本科', 'evidence_ids': [True]},
])
def test_unverified_hard_conditions_are_not_used(tmp_path, model, entry):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3 负责客户开发 具备沟通能力\n基本学历 | 本科', encoding='utf-8')
    model['hard_conditions'] = [entry]
    result = ai.extract_job_with_ai(path)
    assert result['hard_conditions'] == []
    assert '硬性条件有1项无法对应原文' in result['warning']


@pytest.mark.parametrize('initial_has_degree', [False, True])
def test_dedicated_review_recovers_missing_conditions(tmp_path, model, monkeypatch, initial_has_degree):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3 负责客户开发 具备沟通能力\n基本学历 | 本科\n工作经验 | 3年以上', encoding='utf-8')
    degree = {'text': '基本学历 | 本科', 'evidence_ids': [1]}
    experience = {'text': '工作经验 | 3年以上', 'evidence_ids': [2]}
    model['hard_conditions'] = [degree] if initial_has_degree else []
    def analyze(self, payload, job):
        return {'hard_conditions': [degree, experience]} if self.task == 'job_hard_conditions' else model
    monkeypatch.setattr(ai.ExternalModelAdapter, 'analyze', analyze)
    result = ai.extract_job_with_ai(path)
    assert result['hard_conditions'] == ['基本学历 | 本科', '工作经验 | 3年以上']


def test_review_failure_preserves_initial_extraction_with_warning(tmp_path, model, monkeypatch):
    path = tmp_path / 'job.txt'
    path.write_text('销售专员 营销类 P3 负责客户开发 具备沟通能力\n基本学历 | 本科', encoding='utf-8')
    model['hard_conditions'] = [1]
    def analyze(self, payload, job):
        if self.task == 'job_hard_conditions':
            raise ModelError('连接失败')
        return model
    monkeypatch.setattr(ai.ExternalModelAdapter, 'analyze', analyze)
    result = ai.extract_job_with_ai(path)
    assert result['hard_conditions'] == ['基本学历 | 本科']
    assert result['responsibilities'] == ['负责客户开发']
    assert '专项核对未完成' in result['warning']
