from __future__ import annotations
from . import helpers
from .helpers import update_candidate_fields, get_talent_pool, create_task_from_talent_pool, create_fit_task_from_talent_pool, create_job_from_description

import json
import os
import sqlite3
from pathlib import Path

import fitz
import pytest
from docx import Document

from matching_core.config import DISCLAIMER
from .fixture_data import DEMO_DIR
from matching_core.db import rows
from .fixture_data import init_db
from matching_core.services.parser import extract_pdf, parse_candidate
from matching_core.services.privacy import contains_unmasked_sensitive, redact_sensitive
from matching_core.services.report import export_fit_report, export_report
from matching_core.services.scoring import calculate_score, validate_weights
from matching_core.services.tagger import build_tags
from matching_core.services.workflow import (
    archive_job,
    create_fit_task,
    create_job,
    create_task,
    get_candidate,
    get_candidate_job_matches,
    get_job,
    get_jobs,
    get_ranked_job_matches,
    get_sorted_candidates,
    process_fit_task,
    process_task,
    rescore_candidate,
    restore_job,
    save_fit_decision,
    save_review,
    update_job_profile,
    update_job_weights,
)


@pytest.fixture()
def db_path(tmp_path: Path) -> Path:
    path = tmp_path / "test.db"
    init_db(path)
    return path


def sample_resume(job_id: str = "JOB-TECH-001", index: int = 0) -> Path:
    return sorted((DEMO_DIR / "resumes" / job_id).glob("*.pdf"))[index]


def test_database_auto_initializes_and_jobs_load(db_path: Path):
    jobs = get_jobs(db_path)
    assert len(jobs) >= 3
    assert {"高级后端开发工程师", "大客户销售经理", "招聘专员"}.issubset({j["name"] for j in jobs})
    assert all(sum(j["weights"].values()) == 100 for j in jobs)


@pytest.mark.usefixtures("offline_job_extraction")
def test_job_library_supports_manual_add_edit_delete_restore_and_pdf_import(db_path: Path, tmp_path: Path):
    template = get_job("JOB-HR-001", db_path)
    copied_id = create_job(
        "包装质量管理员（测试）", "生产技术序列", "P4", "JD-TEST-1.0",
        template["responsibilities"], template["requirements"], template["hard_conditions"],
        template["dimensions"], template["weights"], "测试岗位", "招聘人员", db_path,
    )
    copied = get_job(copied_id, db_path)
    assert copied["name"] == "包装质量管理员（测试）"
    assert copied["dimensions"] == template["dimensions"]

    update_job_profile(
        copied_id, "包装质量主管（测试）", "生产技术序列", "P5", "JD-TEST-2.0",
        ["负责包装质量检查"], ["具备质量管理经验"], ["2年以上相关经验"],
        "测试岗位", "招聘人员", db_path,
    )
    assert get_job(copied_id, db_path)["level"] == "P5"
    archive_job(copied_id, "招聘人员", db_path)
    assert copied_id not in {item["id"] for item in get_jobs(db_path)}
    assert get_job(copied_id, db_path)["is_active"] == 0
    restore_job(copied_id, "招聘人员", db_path)
    assert copied_id in {item["id"] for item in get_jobs(db_path)}

    pdf_path = tmp_path / "job.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text(
            (40, 40), "岗位职责\n负责数字化运营分析\n任职要求\n具备运营分析经验", fontname="china-s",
        )
        document.save(pdf_path)
    imported_id = create_job_from_description(
        pdf_path, template["id"], "招聘人员", "数字化运营专员（测试）", "职能序列", "P4", db_path,
    )
    imported = get_job(imported_id, db_path)
    assert imported["name"] == "数字化运营专员（测试）"
    assert imported["responsibilities"] == ["负责数字化运营分析"]
    assert imported["requirements"] == ["具备运营分析经验"]


def test_manual_job_supports_custom_dimensions_and_extra_sections(db_path: Path):
    dimensions = [
        {"name": "制造现场管理", "keywords": ["制造", "现场管理"], "negative_keywords": []},
        {"name": "安全与质量", "keywords": ["安全", "质量"], "negative_keywords": []},
    ]
    job_id = create_job(
        "自定义岗位（测试）", "生产序列", "P4", "JD-CUSTOM-1.0",
        ["负责生产现场协调"], ["具备现场管理经验"], ["持有相关资格证书"],
        dimensions, {"制造现场管理": 60.0, "安全与质量": 40.0},
        "测试范围", "招聘人员", db_path,
        extra_sections={"工作地点": "示例园区", "班制说明": "根据生产计划排班"},
    )
    created = get_job(job_id, db_path)
    assert created["dimensions"] == dimensions
    assert created["weights"] == {"制造现场管理": 60.0, "安全与质量": 40.0}
    assert created["extra_sections"]["工作地点"] == "示例园区"

    update_job_profile(
        job_id, created["name"], created["sequence"], created["level"], created["jd_version"],
        created["responsibilities"], created["requirements"], [item["name"] for item in created["hard_conditions"]],
        created["scope"], "招聘人员", db_path,
        extra_sections={"汇报关系": "向生产负责人汇报"},
    )
    assert get_job(job_id, db_path)["extra_sections"] == {"汇报关系": "向生产负责人汇报"}


def test_demo_assets_complete_and_labeled():
    resumes = list((DEMO_DIR / "resumes").rglob("*.pdf"))
    minutes = list((DEMO_DIR / "interview_minutes").rglob("*.txt"))
    jds = list((DEMO_DIR / "job_descriptions").glob("*.docx"))
    assert len(resumes) == 15
    assert len(minutes) >= 6
    assert len(jds) == 3
    assert len(list((DEMO_DIR / "competency_models").glob("*.json"))) == 3
    assert len(list((DEMO_DIR / "capability_dictionaries").glob("*.json"))) == 3
    assert len(list((DEMO_DIR / "scoring_weights").glob("*.json"))) == 3
    for path in minutes:
        assert "仅用于系统演示，内容纯属虚构" in path.read_text(encoding="utf-8")



def test_generated_pdfs_are_text_extractable_and_labeled():
    for path in (DEMO_DIR / "resumes").rglob("*.pdf"):
        result = extract_pdf(path)
        assert result.status == "解析成功"
        assert "候选人编号" in result.text
        assert "仅用于系统演示，内容纯属虚构" in result.text


def test_phone_and_id_redaction_exact_rules():
    result = redact_sensitive("手机13812345678，身份证11010119900101123X")
    assert result.detected
    assert result.text == "手机138****5678，身份证110101********123X"
    assert result.types == ["中国大陆11位手机号", "18位中国大陆身份证号"]
    assert not contains_unmasked_sensitive(result.text)


def test_missing_fields_are_pending_confirmation():
    parsed = parse_candidate("候选人编号：C-1\n姓名：林甲\n专业技能：Python", "a.pdf")
    assert parsed["full_name"] == "林甲"
    assert parsed["masked_name"] == "林*"
    assert parsed["fields"]["education"]["value"] == "待人工确认"
    assert parsed["fields"]["education"]["status"] == "待人工确认"
    assert parsed["parse_status"] == "部分字段待人工确认"
    assert parsed["fields"]["skills"]["location"] == "第3行"


def test_scanned_or_empty_pdf_is_graceful(tmp_path: Path):
    path = tmp_path / "blank.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    result = extract_pdf(path)
    assert result.status == "无法可靠解析，请人工确认"
    assert "空白页" in result.warning


def test_tags_only_claim_evidence_and_mark_absence(db_path: Path):
    job = get_job("JOB-TECH-001", db_path)
    tags = build_tags("候选人编号：C-1\n专业技能：Python、FastAPI", job)
    assert any(t["score"] > 0 and "Python" in t["evidence"] for t in tags)
    assert any(t["score"] > 0 and t["location"] == "第2行" for t in tags)
    assert all("关键词所在段落" not in t["location"] for t in tags)
    absent = [t for t in tags if t["score"] == 0]
    assert absent and all(t["name"] == "无充分证据" and t["needs_confirmation"] for t in absent)


def test_weighted_score_formula_and_range(db_path: Path):
    job = get_job("JOB-TECH-001", db_path)
    text = "Python Java FastAPI Spring Boot 微服务 高并发 架构设计 Kafka Redis 分布式 MySQL SQL优化 性能优化 缓存 单元测试 CI/CD 可观测性 SLA 故障复盘 技术评审 带领 跨团队 指导"
    tags = build_tags(text, job)
    score = calculate_score(tags, job, text)
    expected = round(sum(x["score"] * x["weight"] / 100 for x in score["dimensions"]), 1)
    assert score["total"] == expected
    assert 0 <= score["total"] <= 100


def test_weight_validation():
    validate_weights({"A": 30, "B": 70})
    with pytest.raises(ValueError, match="100"):
        validate_weights({"A": 30, "B": 60})


def test_multi_pdf_task_parses_redacts_scores_and_sorts(db_path: Path):
    from matching_core.services.scoring import score_sort_key
    paths = [sample_resume(index=0), sample_resume(index=1)]
    task_id = create_task("JOB-TECH-001", [p.name for p in paths], "招聘人员", db_path)
    ids = process_task(task_id, "JOB-TECH-001", paths, "招聘人员", db_path=db_path)
    assert len(ids) == 2
    candidates = get_sorted_candidates(task_id, db_path)
    assert len(candidates) == 2
    assert candidates == sorted(candidates, key=lambda item: score_sort_key(item["score"]), reverse=True)
    for candidate in candidates:
        redacted = Path(candidate["redacted_path"]).read_text(encoding="utf-8")
        assert not contains_unmasked_sensitive(redacted)
        assert candidate["tags"]
        assert candidate["score"]["dimensions"]


def test_people_to_jobs_flow_scores_every_selected_job(db_path: Path):
    from matching_core.services.scoring import score_sort_key
    paths = [sample_resume("JOB-TECH-001", 0), sample_resume("JOB-SALES-001", 0)]
    job_ids = [job["id"] for job in get_jobs(db_path)]
    task_id = create_fit_task(job_ids, [path.name for path in paths], "招聘人员", db_path)
    candidate_ids = process_fit_task(task_id, paths, "招聘人员", db_path=db_path)
    assert len(candidate_ids) == 2
    matches = get_candidate_job_matches(task_id=task_id, db_path=db_path)
    assert len(matches) == len(candidate_ids) * len(job_ids)
    for candidate_id in candidate_ids:
        ranked = get_ranked_job_matches(candidate_id, db_path)
        assert {item["job_id"] for item in ranked} == set(job_ids)
        assert ranked == sorted(ranked, key=lambda item: score_sort_key(item["score"]), reverse=True)
        assert all(0 <= item["evidence_completeness"] <= 100 for item in ranked)
        assert all(item["confidence"] in {"高", "中", "低"} for item in ranked)
        assert all(item["recommendation_level"] for item in ranked)

    best = get_ranked_job_matches(candidate_ids[0], db_path)[0]
    save_fit_decision(candidate_ids[0], best["job_id"], "建议业务部门复核关键项目证据", "推荐至该岗位", "招聘人员", db_path)
    updated = get_ranked_job_matches(candidate_ids[0], db_path)[0]
    assert updated["decision"] == "推荐至该岗位"
    assert "复核" in updated["review_comment"]
    with pytest.raises(ValueError):
        save_fit_decision(candidate_ids[0], best["job_id"], "", "自动调岗", "招聘人员", db_path)


def test_resume_is_uploaded_once_and_reused_between_both_flows(db_path: Path):
    source_path = sample_resume("JOB-TECH-001", 0)
    source_task = create_task("JOB-TECH-001", [source_path.name], "招聘人员", db_path)
    source_id = process_task(source_task, "JOB-TECH-001", [source_path], "招聘人员", db_path=db_path)[0]
    source = get_candidate(source_id, db_path)
    job_ids = [job["id"] for job in get_jobs(db_path)]

    fit_task, fit_ids = create_fit_task_from_talent_pool(job_ids, [source_id], "招聘人员", db_path)
    assert len(fit_ids) == 1
    fit_candidate = get_candidate(fit_ids[0], db_path)
    assert fit_candidate["source_candidate_id"] == source_id
    assert fit_candidate["original_path"] == source["original_path"]
    assert fit_candidate["redacted_path"] == source["redacted_path"]
    assert fit_candidate["structured_path"] != source["structured_path"]
    assert len(get_candidate_job_matches(task_id=fit_task, db_path=db_path)) == len(job_ids)

    second_task, second_ids = create_task_from_talent_pool("JOB-SALES-001", [source_id], "招聘人员", db_path)
    assert len(second_ids) == 1
    second_candidate = get_candidate(second_ids[0], db_path)
    assert second_candidate["source_candidate_id"] == source_id
    assert second_candidate["original_path"] == source["original_path"]
    assert second_candidate["redacted_path"] == source["redacted_path"]
    assert second_candidate["job_id"] == "JOB-SALES-001"
    assert second_candidate["task_id"] == second_task

    pool_ids = {candidate["id"] for candidate in get_talent_pool(db_path)}
    assert source_id in pool_ids
    assert fit_ids[0] not in pool_ids
    assert second_ids[0] not in pool_ids
    operations = {item["operation_type"] for item in rows("SELECT * FROM audits", db_path=db_path)}
    assert {"人员库选用", "人员材料复用"}.issubset(operations)


def processed_candidate(db_path: Path) -> tuple[str, str]:
    path = sample_resume()
    task_id = create_task("JOB-TECH-001", [path.name], "招聘人员", db_path)
    candidate_id = process_task(task_id, "JOB-TECH-001", [path], "招聘人员", db_path=db_path)[0]
    return task_id, candidate_id


def test_field_correction_and_confirmation_are_saved(db_path: Path):
    _, candidate_id = processed_candidate(db_path)
    update_candidate_fields(candidate_id, {"education": "人工确认：虚构大学本科"}, "招聘人员", db_path)
    candidate = get_candidate(candidate_id, db_path)
    assert candidate["parsed"]["fields"]["education"]["status"] == "人工已确认"
    assert candidate["parse_status"] == "人工已复核"


def test_manual_decision_defaults_empty_and_accepts_only_fixed_values(db_path: Path):
    _, candidate_id = processed_candidate(db_path)
    assert get_candidate(candidate_id, db_path)["decision"] is None
    save_review(candidate_id, "证据已核验", "进入下一轮", "招聘人员", db_path)
    assert get_candidate(candidate_id, db_path)["decision"] == "进入下一轮"
    with pytest.raises(ValueError):
        save_review(candidate_id, "", "自动录用", "招聘人员", db_path)


def test_weight_change_and_rescore(db_path: Path):
    _, candidate_id = processed_candidate(db_path)
    before = get_candidate(candidate_id, db_path)["score"]["total"]
    weights = {"后端工程能力": 10, "系统架构能力": 10, "数据与性能": 10, "交付与质量": 20, "协作与领导": 50}
    update_job_weights("JOB-TECH-001", weights, "用人部门", db_path)
    after = rescore_candidate(candidate_id, "用人部门", db_path)["total"]
    assert get_job("JOB-TECH-001", db_path)["weights"] == weights
    # No configured scoring model in offline tests: neither run fabricates 0.
    assert before is None and after is None
    score = get_candidate(candidate_id, db_path)["score"]
    assert {d['dimension']: d['weight'] for d in score['dimensions']} == weights


def test_report_is_editable_complete_and_contains_no_sensitive_values(db_path: Path, tmp_path: Path, monkeypatch):
    from .report_fixtures import configure_report_assessment
    configure_report_assessment(monkeypatch)
    _, candidate_id = processed_candidate(db_path)
    candidate = get_candidate(candidate_id, db_path)
    job = get_job(candidate["job_id"], db_path)
    assert candidate["full_name"] == candidate["parsed"]["full_name"]
    assert "*" not in candidate["full_name"]
    confirmed_education = "人工确认：虚构大学本科（复核修正版）"
    update_candidate_fields(candidate_id, {"education": confirmed_education}, "招聘人员", db_path)
    rescore_candidate(candidate_id, "招聘人员", db_path)
    save_review(candidate_id, "已核对原始材料并确认字段修正", "进入下一轮", "招聘人员", db_path)
    # Deliberately pass the stale pre-review object: export must reload the latest confirmed version.
    path = export_report(candidate, job, "招聘人员", db_path, tmp_path / "report.docx")
    assert path.exists() and path.stat().st_size > 0
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    for required in ("人岗匹配评估报告", "综合匹配度", "各维度评分", "硬性条件检查", "关键证据", "关键优势", "主要短板", "风险提示", "AI辅助建议", "人工复核", "人工最终决策", "生成时间", DISCLAIMER):
        assert required in text
    assert confirmed_education in text
    assert "已核对原始材料并确认字段修正" in text
    assert "进入下一轮" in text
    assert "姓名" in text and candidate["full_name"] in text
    assert candidate["masked_name"] not in text
    header_footer = "\n".join(
        paragraph.text
        for section in doc.sections
        for container in (section.header, section.footer)
        for paragraph in container.paragraphs
    )
    complete_text = text + "\n" + header_footer
    for removed in ("星云科技", "可编辑 Word 报告", "AI辅助结果 + 人工最终决策", "仅用于系统演示", "内容纯属虚构"):
        assert removed not in complete_text
    assert len(doc.inline_shapes) == 0
    assert not contains_unmasked_sensitive(text)


def test_report_export_requires_confirmed_fields_and_final_decision(db_path: Path, tmp_path: Path):
    _, candidate_id = processed_candidate(db_path)
    candidate = get_candidate(candidate_id, db_path)
    job = get_job(candidate["job_id"], db_path)
    with pytest.raises(ValueError, match="报告生成条件未满足"):
        export_report(candidate, job, "招聘人员", db_path, tmp_path / "blocked.docx")
    assert not (tmp_path / "blocked.docx").exists()
    assert rows("SELECT * FROM reports WHERE candidate_id=?", (candidate_id,), db_path) == []


def test_fit_report_contains_cross_job_ranking_and_manual_confirmation(db_path: Path, tmp_path: Path):
    source = sample_resume("JOB-TECH-001", 0)
    job_list = get_jobs(db_path)
    task_id = create_fit_task([job["id"] for job in job_list], [source.name], "招聘人员", db_path)
    candidate_id = process_fit_task(task_id, [source], "招聘人员", db_path=db_path)[0]
    matches = get_ranked_job_matches(candidate_id, db_path)
    best = matches[0]
    save_fit_decision(candidate_id, best["job_id"], "已核验核心项目", "推荐至该岗位", "招聘人员", db_path)
    matches = get_ranked_job_matches(candidate_id, db_path)
    candidate = get_candidate(candidate_id, db_path)
    jobs = {job["id"]: job for job in job_list}
    path = export_fit_report(candidate, matches, jobs, "招聘人员", db_path, tmp_path / "fit-report.docx")
    assert path.exists() and path.stat().st_size > 0
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    for required in ("人才适岗推荐报告", "岗位适配排序", "材料支持范围", "匹配结论", "人工推荐结论", "推荐至该岗位", DISCLAIMER):
        assert required in text
    assert all(job["name"] in text for job in job_list)
    assert not contains_unmasked_sensitive(text)


def test_audit_records_required_operations(db_path: Path, tmp_path: Path, monkeypatch):
    from .report_fixtures import configure_report_assessment
    configure_report_assessment(monkeypatch)
    _, candidate_id = processed_candidate(db_path)
    update_candidate_fields(candidate_id, {"skills": "Python（人工确认）"}, "招聘人员", db_path)
    rescore_candidate(candidate_id, "招聘人员", db_path)
    save_review(candidate_id, "已核验", "保留", "招聘人员", db_path)
    candidate = get_candidate(candidate_id, db_path)
    export_report(candidate, get_job(candidate["job_id"], db_path), "招聘人员", db_path, tmp_path / "audit-report.docx")
    operations = {x["operation_type"] for x in rows("SELECT * FROM audits", db_path=db_path)}
    assert {"文件上传", "脱敏处理", "字段修正", "人工决策", "报告导出"}.issubset(operations)


