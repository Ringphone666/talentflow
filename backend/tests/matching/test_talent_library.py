import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from matching_core import db
from matching_core.services import talent_library as library


@pytest.fixture
def database(tmp_path):
    path = tmp_path / "library.db"
    db.init_db(path)
    return path


@pytest.fixture
def local_parser(monkeypatch):
    def parse(path, allow_ocr=False):
        text = Path(path).read_text(encoding="utf-8")
        if "坏文件" in text:
            raise ValueError("简历内容无法读取")
        return ({"full_name": "测试人员", "masked_name": "测*员", "parse_status": "部分字段待人工确认",
                 "fields": {"education": {"value": "本科", "status": "已识别"}},
                 "raw_evidence": [{"file": Path(path).name, "location": "第1行", "text": text}]},
                {"detected": False, "types": []}, text)
    monkeypatch.setattr(library, "parse_material", parse)
    monkeypatch.setattr(library, "enrich_candidate", lambda parsed: parsed)


def write_file(tmp_path, name, text="本科，负责项目管理与协调工作"):
    path = tmp_path / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def test_independent_batch_import_keeps_distinct_same_name_sources(database, tmp_path, local_parser):
    batch = library.create_batch(" 九月校招一批 ", "校招", "招聘人员", database)
    left = write_file(tmp_path / "left", "简历.txt")
    right = write_file(tmp_path / "right", "简历.txt", "本科，负责财务工作")
    result = library.import_files(batch, [left, right], "招聘人员", database)
    assert len(result["imported"]) == 2 and not result["failed"]
    entries = library.list_entries(batch, database)
    assert len({item["original_path"] for item in entries}) == 2
    assert {Path(e["original_path"]).read_text(encoding="utf-8") for e in entries} == {left.read_text(encoding="utf-8"), right.read_text(encoding="utf-8")}
    assert all(Path(e["original_path"]).is_relative_to(tmp_path / "talent") for e in entries)
    assert all("score" not in e and "score" not in json.loads(Path(e["structured_path"]).read_text(encoding="utf-8")) for e in entries)
    assert library.get_batch(batch, database)["parsed_count"] == 2
    assert not db.rows("SELECT * FROM jobs", db_path=database)
    assert not db.rows("SELECT * FROM tasks", db_path=database)


def test_failed_file_kept_and_retry_does_not_duplicate_or_block_other_people(database, tmp_path, local_parser):
    batch = library.create_batch("社招第一批", "社招", "招聘人员", database)
    good, bad = write_file(tmp_path, "成功.txt"), write_file(tmp_path, "失败.txt", "坏文件")
    result = library.import_files(batch, [bad, good], "招聘人员", database)
    assert len(result["imported"]) == len(result["failed"]) == 1
    failed = result["failed"][0]
    assert failed["status"] == "解析失败" and Path(failed["original_path"]).exists()
    assert library.get_batch(batch, database)["failed_count"] == 1
    Path(failed["original_path"]).write_text("已修复，本科学历", encoding="utf-8")
    retried = library.retry_entry(failed["id"], "招聘人员", database)
    assert retried["status"] == "已解析" and not retried["error_message"]
    assert library.get_batch(batch, database)["parsed_count"] == 2
    assert library.get_batch(batch, database)["failed_count"] == 0


def test_real_text_parser_stores_facts_and_redacts_contact_information(database, tmp_path, monkeypatch):
    monkeypatch.setattr(library, "enrich_candidate", lambda parsed: parsed)
    batch = library.create_batch("资料入库", "校招", "招聘人员", database)
    source = write_file(tmp_path, "人员.txt", "姓名：张三\n联系电话：13812345678\n教育背景：本科，信息管理专业\n工作经历：三年项目管理工作")
    result = library.import_files(batch, [source], "招聘人员", database)
    assert not result["failed"]
    entry = result["imported"][0]
    assert entry["sensitive_detected"] == 1
    assert "13812345678" not in Path(entry["redacted_path"]).read_text(encoding="utf-8")
    assert "13812345678" in Path(entry["original_path"]).read_text(encoding="utf-8")


@pytest.mark.parametrize("category", ["全职", "", "所有分类"])
def test_categories_are_validated(database, category):
    with pytest.raises(ValueError):
        library.create_batch("批次", category, "招聘人员", database)


def test_permissions_are_enforced_before_writes(database, tmp_path):
    with pytest.raises(PermissionError):
        library.create_batch("非法", "社招", "用人部门", database)
    batch = library.create_batch("正式", "社招", "招聘人员", database)
    with pytest.raises(PermissionError):
        library.import_files(batch, [write_file(tmp_path, "简历.txt")], "面试官", database)
    assert library.list_entries(batch, database) == []
    assert not (tmp_path / "talent").exists()


def test_classification_and_batch_name_can_be_updated(database):
    batch = library.create_batch("历史导入", "未分类", "招聘人员", database)
    updated = library.update_batch(batch, "九月人才储备", "社招", "招聘人员", database)
    assert (updated["name"], updated["category"]) == ("九月人才储备", "社招")


def doc_text(path):
    document = Document(path)
    return "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])


def test_single_and_multiple_exports_preserve_selection_and_source_files(database, tmp_path, local_parser):
    batch = library.create_batch("校招", "校招", "招聘人员", database)
    paths = [write_file(tmp_path / str(i), "同名简历.txt", f"人员{i}本科") for i in range(2)]
    entries = library.import_files(batch, paths, "招聘人员", database)["imported"]
    original = library.export_resumes(batch, [entries[0]["id"]], "用人部门", database)
    assert original.read_bytes() == paths[0].read_bytes()
    zipped = library.export_resumes(batch, [entry["id"] for entry in entries], "招聘人员", database)
    with ZipFile(zipped) as archive:
        assert len(set(archive.namelist())) == 2
        assert {archive.read(name) for name in archive.namelist()} == {path.read_bytes() for path in paths}
    report = library.export_profiles(batch, [entries[0]["id"]], "招聘人员", database)
    text = doc_text(report)
    assert "人员简历资料报告" in text and "本科" in text
    assert "匹配度" not in text and "录用" not in text.replace("岗位匹配和录用结论请查看对应匹配任务报告", "")
    reports = library.export_profiles(batch, [entry["id"] for entry in entries], "招聘人员", database)
    with ZipFile(reports) as archive:
        assert len(set(archive.namelist())) == 2
        assert all(name.endswith(".docx") for name in archive.namelist())


@pytest.mark.parametrize("export", [library.export_resumes, library.export_profiles])
def test_exports_reject_cross_batch_selection_and_unauthorized_roles(database, tmp_path, local_parser, export):
    one = library.create_batch("第一批", "社招", "招聘人员", database)
    two = library.create_batch("第二批", "校招", "招聘人员", database)
    entry = library.import_files(two, [write_file(tmp_path, "人员.txt")], "招聘人员", database)["imported"][0]
    with pytest.raises(ValueError):
        export(one, [entry["id"]], "招聘人员", database)
    with pytest.raises(PermissionError):
        export(two, [entry["id"]], "面试官", database)
    assert not (tmp_path / "talent" / "exports").exists()


def test_missing_original_is_not_silently_omitted_from_selected_export(database, tmp_path, local_parser):
    batch = library.create_batch("资料", "社招", "招聘人员", database)
    entries = library.import_files(batch, [write_file(tmp_path, "甲.txt"), write_file(tmp_path, "乙.txt")], "招聘人员", database)["imported"]
    Path(entries[1]["original_path"]).unlink()
    with pytest.raises(ValueError, match="原始文件缺失"):
        library.export_resumes(batch, [entry["id"] for entry in entries], "招聘人员", database)


@pytest.mark.parametrize("redacted_problem", [None, "missing", "empty"])
@pytest.mark.parametrize("batch_name", ["历史社招", ""])
def test_legacy_migration_is_idempotent_offline_and_materials_survive_task_deletion(tmp_path, monkeypatch, redacted_problem, batch_name):
    from .fixture_data import init_db
    from matching_core.services import workflow
    database = tmp_path / "legacy.db"
    init_db(database)
    source = write_file(tmp_path, "旧简历.txt", "姓名：测试人员\n教育背景：本科\n工作经历：负责协调工作")
    job = workflow.get_jobs(database)[0]
    task = workflow.create_task(job["id"], [source.name], "招聘人员", database, batch_name=batch_name)
    with db.connect(database) as conn:
        conn.execute("UPDATE tasks SET created_at='2026-09-28T12:34:56+08:00' WHERE id=?", (task,))
    cid = workflow.process_resume(task, job["id"], source, "招聘人员", db_path=database)
    old = workflow.get_candidate(cid, database)
    if redacted_problem == "missing":
        Path(old["redacted_path"]).unlink()
    elif redacted_problem == "empty":
        Path(old["redacted_path"]).write_text("   ", encoding="utf-8")
    with db.connect(database) as conn:
        conn.execute("""INSERT INTO candidates(id,task_id,job_id,source_candidate_id,full_name,masked_name,source_filename,
            original_path,redacted_path,structured_path,redaction_status,parse_status,parsed_json,created_at,updated_at)
            SELECT 'CLONE',task_id,job_id,id,full_name,masked_name,source_filename,original_path,redacted_path,structured_path,
            redaction_status,parse_status,parsed_json,created_at,updated_at FROM candidates WHERE id=?""", (cid,))
    monkeypatch.setattr(library, "parse_material", lambda *a, **kw: pytest.fail("迁移不能重新解析或外发"))
    monkeypatch.setattr(library, "enrich_candidate", lambda *a, **kw: pytest.fail("迁移不能调用模型"))
    assert library.ensure_legacy_talent(database) == {"batches": 1, "entries": 1}
    assert library.ensure_legacy_talent(database) == {"batches": 0, "entries": 0}
    entries = library.list_entries(db_path=database)
    assert len(entries) == 1 and entries[0]["legacy_candidate_id"] == cid
    assert entries[0]["status"] == ("解析失败" if redacted_problem else "已解析")
    assert entries[0]["original_path"] != old["original_path"]
    expected_name = batch_name or f"2026-09-28 12:34 · {job['name']} · 历史导入"
    assert library.get_batch(entries[0]["batch_id"], database)["name"] == expected_name
    assert db.row("SELECT talent_batch_id FROM tasks WHERE id=?", (task,), database)["talent_batch_id"] == entries[0]["batch_id"]
    Path(old["original_path"]).unlink()
    with db.connect(database) as conn:
        conn.execute("DELETE FROM candidates WHERE task_id=?", (task,))
        conn.execute("DELETE FROM tasks WHERE id=?", (task,))
    assert Path(entries[0]["original_path"]).read_bytes() == source.read_bytes()
    assert library.get_entry(entries[0]["id"], database)["parsed"]["raw_evidence"]


def test_task_schema_migration_adds_talent_source_without_resetting_tasks(database):
    with db.connect(database) as conn:
        conn.execute("ALTER TABLE tasks DROP COLUMN talent_batch_id")
    db.init_db(database)
    with db.connect(database) as conn:
        assert "talent_batch_id" in {item["name"] for item in conn.execute("PRAGMA table_info(tasks)")}


def test_delete_batches_physically_removes_selected_records_and_unreferenced_files(database, tmp_path, local_parser):
    one = library.create_batch("第一批", "社招", "招聘人员", database)
    two = library.create_batch("第二批", "校招", "招聘人员", database)
    entries = []
    for batch, name in ((one, "甲.txt"), (two, "乙.txt")):
        entries.extend(library.import_files(batch, [write_file(tmp_path, name)], "招聘人员", database)["imported"])
    assert library.delete_batches([one, two, one], "招聘人员", database) == {"deleted_batches": 2, "deleted_entries": 2}
    assert not library.list_batches(database) and not library.list_entries(db_path=database)
    for entry in entries:
        assert all(not Path(entry[field]).exists() for field in ("original_path", "redacted_path", "structured_path"))
    assert (tmp_path / "甲.txt").exists() and (tmp_path / "乙.txt").exists()
    audits = db.rows("SELECT * FROM audits WHERE operation_type='删除人员批次'", db_path=database)
    assert len(audits) == 2 and {item["object_id"] for item in audits} == {one, two}


def test_delete_batches_permissions_and_all_or_nothing_validation(database, tmp_path, local_parser):
    batch = library.create_batch("正式", "社招", "招聘人员", database)
    entry = library.import_files(batch, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    with pytest.raises(PermissionError, match="无权删除"):
        library.delete_batches([batch], "用人部门", database)
    for selection in ([batch, "不存在"], [], batch):
        with pytest.raises(ValueError):
            library.delete_batches(selection, "招聘人员", database)
    assert library.get_batch(batch, database)["total_count"] == 1
    assert Path(entry["original_path"]).is_file()
    assert not db.rows("SELECT * FROM audits WHERE operation_type='删除人员批次'", db_path=database)


@pytest.mark.parametrize("task_status", ["待处理", "处理中"])
def test_delete_rejects_whole_selection_when_any_matching_task_is_active(database, tmp_path, local_parser, task_status):
    from .fixture_data import init_db
    from matching_core.services import workflow
    init_db(database)
    one = library.create_batch("空闲批次", "社招", "招聘人员", database)
    two = library.create_batch("匹配中批次", "社招", "招聘人员", database)
    entry = library.import_files(two, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    task = workflow.prepare_talent_task([workflow.get_jobs(database)[0]["id"]], [entry["id"]], two, "招聘人员", database)
    with db.connect(database) as conn:
        conn.execute("UPDATE tasks SET status=? WHERE id=?", (task_status, task))
    with pytest.raises(ValueError, match="匹配任务正在处理中"):
        library.delete_batches([one, two], "招聘人员", database)
    assert {item["id"] for item in library.list_batches(database)} == {one, two}
    assert Path(entry["redacted_path"]).is_file()


def test_delete_rejects_whole_selection_during_resume_parsing_or_retry(database, tmp_path, local_parser, monkeypatch):
    one = library.create_batch("空闲批次", "社招", "招聘人员", database)
    two = library.create_batch("解析批次", "社招", "招聘人员", database)
    entry = library.import_files(two, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    parse = library.parse_material
    def parse_while_attempting_delete(*args, **kwargs):
        with pytest.raises(ValueError, match="简历正在解析"):
            library.delete_batches([one, two], "招聘人员", database)
        return parse(*args, **kwargs)
    monkeypatch.setattr(library, "parse_material", parse_while_attempting_delete)
    assert library.retry_entry(entry["id"], "招聘人员", database)["status"] == "已解析"
    assert {item["id"] for item in library.list_batches(database)} == {one, two}


def test_delete_is_blocked_between_files_of_one_import(database, tmp_path, local_parser, monkeypatch):
    batch = library.create_batch("连续导入", "校招", "招聘人员", database)
    files = [write_file(tmp_path, "第一份.txt"), write_file(tmp_path, "第二份.txt")]
    parse_entry = library._parse_entry
    completed = []
    def parse_and_attempt_delete(*args, **kwargs):
        result = parse_entry(*args, **kwargs)
        completed.append(result["id"])
        if len(completed) == 1:
            # The first file is already marked complete; the second must already have a pending record.
            entries = library.list_entries(batch, database)
            assert len(entries) == 2 and sum(item["status"] == "待解析" for item in entries) == 1
            with pytest.raises(ValueError, match="简历正在解析"):
                library.delete_batches([batch], "招聘人员", database)
        return result
    monkeypatch.setattr(library, "_parse_entry", parse_and_attempt_delete)
    result = library.import_files(batch, files, "招聘人员", database)
    assert len(result["imported"]) == 2 and not result["failed"]
    assert library.delete_batches([batch], "招聘人员", database)["deleted_entries"] == 2


def test_delete_library_batch_retains_matching_candidates_reports_and_shared_materials(database, tmp_path, local_parser):
    from .fixture_data import init_db
    from matching_core.services import workflow
    init_db(database)
    batch = library.create_batch("已匹配批次", "社招", "招聘人员", database)
    entry = library.import_files(batch, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    job = workflow.get_jobs(database)[0]["id"]
    task = workflow.prepare_talent_task([job], [entry["id"]], batch, "招聘人员", database)
    cid = workflow._clone_candidate_for_task(entry["id"], task, job, "招聘人员", database)
    report = write_file(tmp_path, "已生成报告.docx", "already exported")
    with db.connect(database) as conn:
        conn.execute("UPDATE tasks SET status='已完成' WHERE id=?", (task,))
        conn.execute("UPDATE candidates SET structured_path=? WHERE id=?", (entry["structured_path"], cid))
        conn.execute("INSERT INTO reports(candidate_id,path,generated_at,operator) VALUES(?,?,?,?)", (cid, str(report), db.now_iso(), "招聘人员"))
    assert library.delete_batches([batch], "招聘人员", database) == {"deleted_batches": 1, "deleted_entries": 1}
    candidate = workflow.get_candidate(cid, database)
    assert all(Path(candidate[field]).is_file() for field in ("original_path", "redacted_path", "structured_path"))
    assert Path(candidate["redacted_path"]).read_text(encoding="utf-8")
    assert db.row("SELECT id FROM tasks WHERE id=?", (task,), database)
    assert len(workflow.get_reports(cid, database)) == 1 and report.is_file()


def test_deletion_never_removes_paths_outside_library_or_used_by_other_entries(database, tmp_path, local_parser):
    one = library.create_batch("待删除", "社招", "招聘人员", database)
    two = library.create_batch("保留", "校招", "招聘人员", database)
    entry = library.import_files(one, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    retained = library.import_files(two, [write_file(tmp_path, "另一人.txt")], "招聘人员", database)["imported"][0]
    outside = write_file(tmp_path, "不属于人才库.txt")
    with db.connect(database) as conn:
        conn.execute("UPDATE talent_entries SET original_path=? WHERE id=?", (entry["original_path"], retained["id"]))
        conn.execute("UPDATE talent_entries SET redacted_path=? WHERE id=?", (str(outside), entry["id"]))
    library.delete_batches([one], "招聘人员", database)
    assert outside.is_file() and Path(entry["original_path"]).is_file()
    assert not Path(entry["structured_path"]).exists()
    assert library.get_batch(two, database)["total_count"] == 1


def test_deleted_legacy_library_batch_does_not_reappear_while_original_candidates_remain(database, tmp_path):
    from .fixture_data import init_db
    from matching_core.services import workflow
    init_db(database)
    source = write_file(tmp_path, "历史简历.txt", "姓名：测试人员\n教育背景：本科\n工作经历：负责协调工作")
    job = workflow.get_jobs(database)[0]["id"]
    task = workflow.create_task(job, [source.name], "招聘人员", database)
    cid = workflow.process_resume(task, job, source, "招聘人员", db_path=database)
    with db.connect(database) as conn:
        conn.execute("UPDATE tasks SET status='已完成' WHERE id=?", (task,))
    library.ensure_legacy_talent(database)
    batch = library.list_batches(database)[0]
    library.delete_batches([batch["id"]], "招聘人员", database)
    assert workflow.get_candidate(cid, database)["parsed"]["raw_evidence"]
    assert library.ensure_legacy_talent(database) == {"batches": 0, "entries": 0}
    assert library.ensure_legacy_talent(database) == {"batches": 0, "entries": 0}
    assert not library.list_batches(database) and not library.list_entries(db_path=database)
    audit = db.row("SELECT * FROM audits WHERE operation_type='删除人员批次' AND object_id=?", (batch["id"],), database)
    assert json.loads(audit["before_json"])["历史任务编号"] == task


def test_delete_audit_and_records_commit_together(database, tmp_path, local_parser):
    import sqlite3
    batch = library.create_batch("待删除", "社招", "招聘人员", database)
    entry = library.import_files(batch, [write_file(tmp_path, "资料.txt")], "招聘人员", database)["imported"][0]
    with db.connect(database) as conn:
        conn.executescript("""CREATE TRIGGER fail_talent_audit BEFORE INSERT ON audits
            WHEN NEW.operation_type='删除人员批次' BEGIN SELECT RAISE(ABORT,'audit unavailable'); END;""")
    with pytest.raises(sqlite3.IntegrityError, match="audit unavailable"):
        library.delete_batches([batch], "招聘人员", database)
    assert library.get_batch(batch, database)["total_count"] == 1
    assert Path(entry["original_path"]).is_file()
