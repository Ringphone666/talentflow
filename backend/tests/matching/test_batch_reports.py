import csv
import io
import json
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document

from matching_core.db import connect, init_db, rows
from matching_core.services import batch_reports as reports


def text_of_docx(value):
    doc = Document(value)
    return "\n".join([*(p.text for p in doc.paragraphs),
                      *(c.text for t in doc.tables for r in t.rows for c in r.cells)])


@pytest.fixture
def batch_db(tmp_path, monkeypatch):
    database = tmp_path / "batch.db"
    monkeypatch.setattr(reports, "REPORT_DIR", tmp_path / "exports")
    init_db(database)
    with connect(database) as conn:
        for i in range(1, 6):
            conn.execute("""INSERT INTO jobs(id,name,sequence,level,jd_version,responsibilities,requirements,
                hard_conditions_json,dimensions_json,weights_json,weight_version,scope,updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
                f"J{i}", f"岗位{i}", "专业", "五级", "V1", "项目交付", "沟通协作", "[]", "[]",
                json.dumps({"项目能力": 100}, ensure_ascii=False), "W1", "测试", "2026-09-28"))
        for task_id, mode, job_ids in (("T", "岗位找人", ["J1"]), ("F", "人找岗位", ["J1", "J2"]), ("OTHER", "岗位找人", ["J1"])):
            conn.execute("""INSERT INTO tasks(id,job_id,task_mode,job_ids_json,status,created_at,updated_at)
                            VALUES(?,?,?,?,?,?,?)""",
                         (task_id, "J1", mode, json.dumps(job_ids), "已完成", "2026-09-28", "2026-09-28"))
    return database


def add_candidate(database, cid, task="T", decision="进入下一轮", reviewed=True, stale=False, failed=False, name="同名人员"):
    score = {
        "scoring_method": "ai_semantic_v1", "dimensions": [
            {"dimension": "项目能力", "weight": 100, "score": 85, "reason": "有项目交付记录",
             "evidence": [{"file": "C:\\private\\简历.pdf", "location": "第1页", "text": "负责项目交付，13812345678"}]}],
        "hard_checks": [], "requires_reassessment": stale,
        "reasoning_analysis": {"status": "调用失败" if failed else "调用成功", "reason": "api_key=super-secret-token"},
    }
    parsed = {"candidate_code": "SAME", "full_name": name,
              "fields": {"phone": {"label": "手机", "value": "13812345678", "status": "人工已确认"}}}
    with connect(database) as conn:
        conn.execute("""INSERT INTO candidates(id,task_id,job_id,full_name,masked_name,source_filename,original_path,
            redacted_path,structured_path,redaction_status,parse_status,parsed_json,score_json,review_comment,decision,created_at,updated_at)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""", (
            cid, task, "J1", name, "同*", "C:\\private\\简历.pdf", "C:\\private\\original.pdf", "", "", "已脱敏",
            "人工已复核" if reviewed else "待人工确认", json.dumps(parsed, ensure_ascii=False),
            json.dumps(score, ensure_ascii=False), "核验项目职责，13812345678，C:\\private\\note.txt，api_key=super-secret-token", decision,
            "2026-09-28", "2026-09-28"))
    return score


def add_match(database, cid, jid, score, decision=None):
    with connect(database) as conn:
        conn.execute("""INSERT INTO candidate_job_matches(task_id,candidate_id,job_id,score_json,review_comment,decision,updated_at)
                        VALUES(?,?,?,?,?,?,?)""", ("F", cid, jid, json.dumps(score, ensure_ascii=False), "已核对岗位相关项目", decision, "2026-09-28"))


@pytest.mark.parametrize("export", [reports.export_batch_reports, reports.export_comparison_report])
def test_enforces_role_and_batch_selection_before_writing(batch_db, export):
    add_candidate(batch_db, "C1")
    add_candidate(batch_db, "C2", task="OTHER")
    with pytest.raises(PermissionError):
        export("T", ["C1"], "面试官", batch_db)
    for selected in ([], ["C2"], ["C1", "C2"], ["deleted"], "C1"):
        with pytest.raises(ValueError):
            export("T", selected, "招聘人员", batch_db)
    assert not reports.REPORT_DIR.exists()
    assert not rows("SELECT * FROM reports", db_path=batch_db)


def test_zip_contains_only_ready_people_and_safe_audited_summary(batch_db):
    add_candidate(batch_db, "C1", name="=SUM(1+1)")
    add_candidate(batch_db, "C2", name="=SUM(1+1)")
    add_candidate(batch_db, "C3", decision=None)
    add_candidate(batch_db, "C4", stale=True)
    add_candidate(batch_db, "C5", failed=True)
    result = reports.export_batch_reports("T", ["C1", "C2", "C3", "C4", "C5", "C1"], "招聘人员", batch_db)
    assert [item["candidate_id"] for item in result["exported"]] == ["C1", "C2"]
    assert {item["candidate_id"] for item in result["blocked"]} == {"C3", "C4", "C5"}
    with ZipFile(result["path"]) as archive:
        names = archive.namelist()
        documents = [name for name in names if name.endswith(".docx")]
        assert len(documents) == len(set(documents)) == 2
        for name in documents:
            text = text_of_docx(io.BytesIO(archive.read(name)))
            assert "138****5678" in text and "简历.pdf" in text
            assert "13812345678" not in text and "private" not in text and "super-secret" not in text
        summary = list(csv.DictReader(io.StringIO(archive.read("审核摘要.csv").decode("utf-8-sig"))))
        assert len(summary) == 5
        assert summary[0]["姓名"].startswith("'=SUM")
        assert summary[2]["人工结论"] == "未确认" and summary[2]["导出结果"] == "未导出"
        assert "private" not in str(summary) and "super-secret" not in str(summary)
    audit = json.loads(rows("SELECT * FROM audits", db_path=batch_db)[-1]["after_json"])
    assert audit["导出类型"] == "单人报告打包"
    assert audit["所选候选人"] == ["C1", "C2", "C3", "C4", "C5"]
    assert len(rows("SELECT * FROM reports", db_path=batch_db)) == 5


def test_fresh_database_review_revocation_blocks_next_export(batch_db):
    add_candidate(batch_db, "C1")
    first = reports.export_batch_reports("T", ["C1"], "招聘人员", batch_db)
    assert first["path"].exists()
    with connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET decision=NULL WHERE id='C1'")
    second = reports.export_batch_reports("T", ["C1"], "招聘人员", batch_db)
    assert second["path"] is None and not second["exported"]
    assert "人工最终决策" in "".join(second["blocked"][0]["reasons"])


def test_selection_comparison_preserves_pending_failed_and_stale_states(batch_db):
    add_candidate(batch_db, "C1", name="已审核人员")
    add_candidate(batch_db, "C2", name="未审核人员", decision=None)
    add_candidate(batch_db, "C3", name="分析失败人员", failed=True)
    add_candidate(batch_db, "C4", name="资料更新人员", stale=True)
    path = reports.export_comparison_report("T", ["C1", "C2", "C3", "C4"], "用人部门", batch_db)
    text = text_of_docx(path)
    for expected in ("同一岗位人员对比", "维度评分对比", "逐人简要分析", "复核意见", "未确认", "待分析（分析失败）", "待重新分析后确认"):
        assert expected in text
    # 已经确认的人不再标“已审核”“已分析”之类的状态。
    for absent in ("审核中", "已分析", "审核依据"):
        assert absent not in text
    assert "关键材料依据" in text and "第1页" in text
    assert "13812345678" not in text and "private" not in text and "super-secret" not in text
    # 第一个表格是结果摘要卡，人员对比表是第二个；各列为人员及编号、综合匹配度、人工结论。
    comparison = Document(path).tables[1]
    # 综合匹配度一列是数据条，表头合并的格子会重复出现。
    assert list(dict.fromkeys(cell.text for cell in comparison.rows[0].cells)) == ["人员及编号", "综合匹配度", "人工结论"]
    assert comparison.rows[1].cells[-1].text == "进入下一轮" and comparison.rows[2].cells[-1].text == "未确认"
    assert comparison.rows[3].cells[1].text == "待分析（分析失败）"
    assert comparison.rows[4].cells[-1].text == "待重新分析后确认"
    assert Document(path).paragraphs[0].style.name == "Title"


def test_fit_zip_requires_confirmed_fields_and_current_complete_job_analysis(batch_db):
    score = add_candidate(batch_db, "F1", task="F", decision=None)
    add_match(batch_db, "F1", "J1", score, "作为备选岗位")
    result = reports.export_batch_reports("F", ["F1"], "招聘人员", batch_db)
    assert result["path"] is None and "部分岗位" in "".join(result["blocked"][0]["reasons"])
    add_match(batch_db, "F1", "J2", score, "推荐至该岗位")
    exported = reports.export_batch_reports("F", ["F1"], "招聘人员", batch_db)
    with ZipFile(exported["path"]) as archive:
        document = next(name for name in archive.namelist() if name.endswith(".docx"))
        text = text_of_docx(io.BytesIO(archive.read(document)))
        assert "人才适岗推荐报告" in text and "推荐至该岗位" in text
    with connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET parse_status='待人工确认' WHERE id='F1'")
    blocked = reports.export_batch_reports("F", ["F1"], "招聘人员", batch_db)
    assert blocked["path"] is None and "字段" in "".join(blocked["blocked"][0]["reasons"])
    with connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET parse_status='人工已复核' WHERE id='F1'")
        stale = deepcopy(score)
        stale["requires_reassessment"] = True
        conn.execute("UPDATE candidate_job_matches SET score_json=? WHERE candidate_id='F1' AND job_id='J2'", (json.dumps(stale),))
    blocked = reports.export_batch_reports("F", ["F1"], "招聘人员", batch_db)
    assert blocked["path"] is None and "重新分析" in "".join(blocked["blocked"][0]["reasons"])


def test_fit_comparison_matrix_contains_missing_pairs_without_inferred_approval(batch_db):
    score = add_candidate(batch_db, "F1", task="F", decision=None, name="已有人岗记录")
    add_candidate(batch_db, "F2", task="F", decision=None, name="待分析人员")
    add_match(batch_db, "F1", "J1", score, "推荐至该岗位")
    add_match(batch_db, "F1", "J2", score)
    path = reports.export_comparison_report("F", ["F1", "F2"], "招聘人员", batch_db)
    # 第一个表格是结果摘要卡，矩阵是第二个；“★”标出优先核验岗位。
    matrix = Document(path).tables[1]
    assert matrix.rows[1].cells[1].text == "★ 85.0\n推荐至该岗位"
    assert matrix.rows[1].cells[2].text == "85.0"
    assert matrix.rows[2].cells[1].text == "待分析"
    text = text_of_docx(path)
    assert "人员与岗位匹配矩阵" in text and "尚无岗位匹配结果" in text
    assert "不代表人工通过" in text


def test_many_roles_split_into_readable_matrix_panels_and_preserve_zero(batch_db):
    score = add_candidate(batch_db, "F1", task="F", decision=None)
    score["dimensions"][0]["score"] = 0
    with connect(batch_db) as conn:
        conn.execute("UPDATE tasks SET job_ids_json=? WHERE id='F'", (json.dumps([f"J{i}" for i in range(1, 6)]),))
    for i in range(1, 6):
        add_match(batch_db, "F1", f"J{i}", score, "暂不推荐")
    path = reports.export_comparison_report("F", ["F1"], "招聘人员", batch_db)
    # 结果摘要卡之后是两组矩阵，再往后是逐人分析的表格。
    tables = Document(path).tables
    assert [len(table.columns) for table in tables[1:3]] == [5, 2]
    # 0 分是有效得分，照常写出，其下是人工结论。
    assert tables[1].rows[1].cells[1].text == "0.0\n暂不推荐"


def test_batch_name_is_used_with_old_task_fallback(batch_db):
    add_candidate(batch_db, "C1")
    fallback = reports.export_batch_reports("T", ["C1"], "招聘人员", batch_db)
    assert "T" in fallback["batch_name"]
    with connect(batch_db) as conn:
        if "batch_name" not in {r["name"] for r in conn.execute("PRAGMA table_info(tasks)")}:
            conn.execute("ALTER TABLE tasks ADD COLUMN batch_name TEXT NOT NULL DEFAULT ''")
        conn.execute("UPDATE tasks SET batch_name='秋季社招第一批' WHERE id='T'")
    named = reports.export_comparison_report("T", ["C1"], "招聘人员", batch_db)
    assert "批次名称：秋季社招第一批" in text_of_docx(named)


@pytest.mark.parametrize("kind", ["zip", "comparison"])
def test_delete_shared_batch_report_preserves_file_until_last_reference(batch_db, monkeypatch, kind):
    from matching_core.services import workflow

    monkeypatch.setattr(workflow, "REPORT_DIR", reports.REPORT_DIR)
    for cid in ("C1", "C2", "C3"):
        add_candidate(batch_db, cid)
    if kind == "zip":
        shared_path = reports.export_batch_reports("T", ["C1", "C2"], "招聘人员", batch_db)["path"]
    else:
        shared_path = reports.export_comparison_report("T", ["C1", "C2"], "招聘人员", batch_db)
    independent_path = reports.export_comparison_report("T", ["C3"], "招聘人员", batch_db)
    shared_rows = rows("SELECT * FROM reports WHERE path=? ORDER BY id", (str(shared_path),), batch_db)
    original_bytes = shared_path.read_bytes()
    assert len(shared_rows) == 2

    workflow.delete_report(shared_rows[0]["id"], "招聘人员", batch_db)
    remaining = rows("SELECT * FROM reports WHERE path=?", (str(shared_path),), batch_db)
    assert [item["id"] for item in remaining] == [shared_rows[1]["id"]]
    assert shared_path.read_bytes() == original_bytes
    assert independent_path.is_file()
    assert len(rows("SELECT * FROM candidates", db_path=batch_db)) == 3

    workflow.delete_report(shared_rows[1]["id"], "招聘人员", batch_db)
    assert not shared_path.exists()
    assert not rows("SELECT * FROM reports WHERE path=?", (str(shared_path),), batch_db)
    assert independent_path.is_file()
    assert len(rows("SELECT * FROM reports", db_path=batch_db)) == 1
    assert len(rows("SELECT * FROM candidates", db_path=batch_db)) == 3
    deleted = rows("SELECT * FROM audits WHERE operation_type='报告删除'", db_path=batch_db)
    assert {item["object_id"] for item in deleted} == {str(item["id"]) for item in shared_rows}


def test_single_export_checks_role_scope_and_rereads_revoked_review(batch_db):
    add_candidate(batch_db, "C1")
    add_candidate(batch_db, "C2", task="OTHER")
    with pytest.raises(PermissionError):
        reports.export_single_report("T", "C1", "面试官", batch_db)
    with pytest.raises(ValueError, match="不属于"):
        reports.export_single_report("T", "C2", "招聘人员", batch_db)
    assert not reports.REPORT_DIR.exists()
    first = reports.export_single_report("T", "C1", "招聘人员", batch_db)
    second = reports.export_single_report("T", "C1", "招聘人员", batch_db)
    assert first != second and first.is_file() and second.is_file()
    text = text_of_docx(first)
    assert "138****5678" in text and "简历.pdf" in text
    assert "13812345678" not in text and "private" not in text and "super-secret" not in text
    with connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET decision=NULL WHERE id='C1'")
    with pytest.raises(ValueError, match="人工最终决策"):
        reports.export_single_report("T", "C1", "招聘人员", batch_db)
    assert len(rows("SELECT * FROM reports", db_path=batch_db)) == 2
    audit = json.loads(rows("SELECT * FROM audits", db_path=batch_db)[-1]["after_json"])
    assert audit["导出类型"] == "单人报告" and audit["所选候选人"] == ["C1"]


@pytest.mark.parametrize("updates,reason", [
    ({"reviewed": False}, "字段"), ({"stale": True}, "重新分析"), ({"failed": True}, "评分尚未完成"),
])
def test_single_export_blocks_unconfirmed_stale_and_failed_analysis(batch_db, updates, reason):
    add_candidate(batch_db, "C1", **updates)
    with pytest.raises(ValueError, match=reason):
        reports.export_single_report("T", "C1", "招聘人员", batch_db)
    assert not reports.REPORT_DIR.exists()
    assert not rows("SELECT * FROM reports", db_path=batch_db)


def test_single_fit_export_requires_complete_jobs_and_latest_field_confirmation(batch_db):
    score = add_candidate(batch_db, "F1", task="F", decision=None)
    add_match(batch_db, "F1", "J1", score, "推荐至该岗位")
    with pytest.raises(ValueError, match="部分岗位"):
        reports.export_single_report("F", "F1", "用人部门", batch_db)
    add_match(batch_db, "F1", "J2", score, "作为备选岗位")
    path = reports.export_single_report("F", "F1", "用人部门", batch_db)
    text = text_of_docx(path)
    assert "人才适岗推荐报告" in text and "岗位1" in text and "岗位2" in text
    with connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET parse_status='待人工确认' WHERE id='F1'")
    with pytest.raises(ValueError, match="字段"):
        reports.export_single_report("F", "F1", "用人部门", batch_db)


def test_csv_sanitization_preserves_supported_zero_and_blocks_private_values():
    values = [reports._csv_cell(value) for value in [
        0, None, "13812345678", "C:\\private\\note.txt", "\ufeff=1+1",
        "https://example.invalid/evidence/project", "http://example.invalid/records",
        '{"api_key": "synthetic-private-value"}',
    ]]
    buffer = io.StringIO(newline="")
    csv.writer(buffer).writerow(values)
    parsed = next(csv.reader(io.StringIO(buffer.getvalue())))
    assert parsed[:5] == ["0", "", "138****5678", "[路径已隐藏]", "'\ufeff=1+1"]
    assert parsed[5:7] == ["https://example.invalid/evidence/project", "http://example.invalid/records"]
    assert "synthetic-private-value" not in parsed[7] and "敏感配置已隐藏" in parsed[7]
