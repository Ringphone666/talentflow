"""Exercise both workflows using a generic synthetic job library."""
from pathlib import Path
from . import helpers

import pytest
from docx import Document

from .fixture_data import DEMO_DIR
from matching_core.db import init_db, rows
from matching_core.services import workflow as w
from matching_core.services.report import export_report, export_fit_report


def test_synthetic_jobs_support_both_flows_and_reviewed_reports(tmp_path, monkeypatch):
    from .report_fixtures import configure_report_assessment
    configure_report_assessment(monkeypatch)
    database = tmp_path / "synthetic.db"
    init_db(database, seed=True)
    jobs = w.get_jobs(database)
    assert len(jobs) == 3
    assert all(job["id"].startswith("JOB-") for job in jobs)
    material = next((DEMO_DIR / "resumes/JOB-HR-001").glob("*.pdf"))
    task = w.create_task(jobs[0]["id"], [material.name], "招聘人员", database)
    candidate_id, = w.process_task(task, jobs[0]["id"], [material], "招聘人员", db_path=database)
    candidate = w.get_candidate(candidate_id, database)
    assert 0 <= candidate["score"]["total"] <= 100
    with pytest.raises(ValueError):
        export_report(candidate, jobs[0], "招聘人员", database, tmp_path / "blocked.docx")
    helpers.update_candidate_fields(candidate_id, {"education": "回归测试：已复核学历材料"}, "招聘人员", database)
    w.rescore_candidate(candidate_id, "招聘人员", database)
    w.save_review(candidate_id, "回归测试：材料已核对", "保留", "招聘人员", database)
    single = export_report(candidate, jobs[0], "招聘人员", database, tmp_path / "single.docx")
    text = "\n".join(cell.text for table in Document(single).tables for row in table.rows for cell in row.cells)
    assert "回归测试：已复核学历材料" in text
    fit_task, fit_ids = helpers.create_fit_task_from_talent_pool([j["id"] for j in jobs], [candidate_id], "招聘人员", database)
    fit_candidate = w.get_candidate(fit_ids[0], database)
    assert fit_candidate["original_path"] == candidate["original_path"]
    matches = w.get_ranked_job_matches(fit_candidate["id"], database)
    assert len(matches) == len(jobs)
    job_map = {job["id"]: job for job in jobs}
    with pytest.raises(ValueError):
        export_fit_report(fit_candidate, matches, job_map, "招聘人员", database, tmp_path / "blocked-fit.docx")
    helpers.update_candidate_fields(fit_candidate["id"], {"education": "回归测试：已复核学历材料"}, "招聘人员", database)
    w.evaluate_candidate_jobs(fit_candidate["id"], [job["id"] for job in jobs], "招聘人员", database)
    w.save_fit_decision(fit_candidate["id"], matches[0]["job_id"], "回归测试：已复核岗位证据", "推荐至该岗位", "招聘人员", database)
    result = export_fit_report(fit_candidate, matches, job_map, "招聘人员", database, tmp_path / "fit.docx")
    assert result.exists() and result.stat().st_size > 0
    exported_document = Document(result)
    exported_text = "\n".join(p.text for p in exported_document.paragraphs) + "\n" + "\n".join(cell.text for table in exported_document.tables for row in table.rows for cell in row.cells)
    assert "回归测试：已复核岗位证据" in exported_text
    with pytest.raises(PermissionError):
        export_fit_report(fit_candidate, matches, job_map, "面试官", database, tmp_path / "denied.docx")
    w.save_fit_decision(fit_candidate["id"], matches[0]["job_id"], "撤销确认", None, "招聘人员", database)
    with pytest.raises(ValueError):
        export_fit_report(fit_candidate, matches, job_map, "招聘人员", database, tmp_path / "revoked.docx")
    assert rows("SELECT status FROM tasks WHERE id=?", (fit_task,), database)[0]["status"] == "已完成"
