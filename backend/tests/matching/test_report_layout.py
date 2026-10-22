"""人岗匹配报告的版式：页眉页脚、摘要卡、数据条、热力着色、证据分组和姓名未确认时的显示。"""
import csv
import io
from copy import deepcopy
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from matching_core.config import DISCLAIMER
from matching_core.services import batch_reports as reports
from matching_core.services import report_style as style
from matching_core.services.report import build_fit_report, build_report

from .test_batch_reports import add_candidate, add_match, batch_db  # noqa: F401
from .test_comprehensive_reports import report_case


def text_of(doc):
    return "\n".join([*(p.text for p in doc.paragraphs),
                      *(cell.text for table in doc.tables for row in table.rows for cell in row.cells)])


def table_headed(doc, first_header):
    return next(table for table in doc.tables if table.rows[0].cells[0].text == first_header)


def bands(row):
    """一行里数据条的颜色：数据条是带底色的空段落。"""
    return [shade.get(qn("w:fill")) for shade in row._tr.iter(qn("w:shd")) if shade.getparent().tag == qn("w:pPr")]


def cell_fill(cell):
    shade = cell._tc.tcPr.find(qn("w:shd"))
    return None if shade is None else shade.get(qn("w:fill"))


def hard_case():
    candidate, job = report_case()
    candidate["score"]["hard_checks"] = [
        {"condition": "本科及以上学历", "status": "满足", "evidence": "学历为本科"},
        {"condition": "三年以上工作经验", "status": "不满足", "evidence": "工作经历不足三年"},
        {"condition": "持有资格证书", "status": "无法确认", "evidence": "材料未提及"},
    ]
    return candidate, job


@pytest.mark.parametrize("name, filename, expected", [
    ("待人工确认", "张三简历.pdf", "姓名待确认（张三简历）"),
    ("待人工确认", "C:\\private\\校招.投递.v2.docx", "姓名待确认（校招.投递.v2）"),
    ("", "uploads/李四 简历.txt", "姓名待确认（李四 简历）"),
    (None, "", "姓名待确认"),
    ("王五", "别人的简历.pdf", "王五"),
])
def test_unconfirmed_name_is_shown_with_resume_file_name(name, filename, expected):
    assert style.person_label(name, filename) == expected


def test_single_report_distinguishes_people_whose_names_are_unconfirmed(tmp_path):
    candidate, job = report_case()
    candidate.update(full_name="待人工确认", masked_name="待人工确认", source_filename="简历_应聘综合事务专员.pdf")
    text = text_of(Document(build_report(candidate, job, tmp_path / "unnamed.docx")))
    assert "候选人姓名：姓名待确认（简历_应聘综合事务专员）" in text
    assert "候选人姓名：待人工确认" not in text
    # 姓名已确认的人员照常显示姓名。
    candidate, job = report_case()
    candidate["source_filename"] = "简历_应聘综合事务专员.pdf"
    text = text_of(Document(build_report(candidate, job, tmp_path / "named.docx")))
    assert "候选人姓名：报告验证人员" in text and "姓名待确认" not in text


def test_page_header_footer_and_page_number_fields(tmp_path):
    candidate, job = report_case()
    doc = Document(build_report(candidate, job, tmp_path / "page.docx"))
    section = doc.sections[0]
    # A4 纵向。
    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    header = "\n".join(p.text for p in section.header.paragraphs)
    assert "TalentFlow · 人力资源智能体" in header and "人岗匹配评估报告" in header
    footer = "\n".join(p.text for p in section.footer.paragraphs)
    assert "第 " in footer and " 页 / 共 " in footer and "AI辅助生成，须人工复核" in footer
    # 页码和总页数是 Word 域，不是写死的数字。
    codes = [node.text.strip() for node in section.footer._element.iter(qn("w:instrText"))]
    assert codes == ["PAGE", "NUMPAGES"]
    kinds = [node.get(qn("w:fldCharType")) for node in section.footer._element.iter(qn("w:fldChar"))]
    assert kinds == ["begin", "separate", "end"] * 2
    fit = Document(build_fit_report(candidate, [{"job_id": job["id"], "score": candidate["score"]}], {job["id"]: job},
                                    tmp_path / "fit.docx"))
    assert "人才适岗推荐报告" in fit.sections[0].header.paragraphs[0].text


def test_summary_card_follows_the_title_with_key_results(tmp_path):
    candidate, job = hard_case()
    doc = Document(build_report(candidate, job, tmp_path / "card.docx"))
    assert doc.paragraphs[0].style.name == "Title"
    card = doc.tables[0]
    # 摘要卡在正文第一个表格之前没有其他表格，且紧跟标题和基本信息。
    body = list(doc.element.body)
    assert body.index(card._tbl) == 2
    text = "\n".join(cell.text for row in card.rows for cell in row.cells)
    for expected in ("综合匹配度：15.0 / 100", "匹配结论：硬性条件存在差距", "覆盖岗位权重", "20%", "人工结论", "保留",
                     "硬性条件核验", "满足 1 项 / 不满足 1 项 / 待核实 1 项"):
        assert expected in text
    assert {cell_fill(cell) for row in card.rows for cell in row.cells} <= {style.LIGHT, style.PALE}
    # 综合匹配度用大号数字。
    sizes = {run.text: run.font.size.pt for cell in card.rows[0].cells for p in cell.paragraphs for run in p.runs}
    assert sizes["15.0"] >= 20 and sizes["综合匹配度："] <= 11


def test_dimension_without_score_has_no_bar_and_is_not_treated_as_zero(tmp_path):
    candidate, job = report_case()
    doc = Document(build_report(candidate, job, tmp_path / "bars.docx"))
    table = table_headed(doc, "岗位评价维度")
    scored, missing = table.rows[1], table.rows[2]
    assert scored.cells[0].text == "沟通协调" and bands(scored) == [style.BAR, style.TRACK]
    assert "75.0" in [cell.text for cell in scored.cells]
    # 着色部分占数据条的比例与得分一致（取 2% 的整数倍）。
    widths = [int(cell.tcPr.find(qn("w:tcW")).get(qn("w:w"))) for cell in scored._tr.iter(qn("w:tc"))]
    filled, rest = widths[1] - style.BAND_INSET, widths[2]
    assert abs(filled / (filled + rest) - 0.76) < 0.005
    assert missing.cells[0].text == "行政统筹" and bands(missing) == []
    texts = [cell.text for cell in missing.cells]
    assert "未形成得分 · 待补充证据" in texts and "未确认贡献" in texts
    assert not any(text.startswith("0") for text in texts)


def test_supported_zero_score_draws_an_empty_bar_instead_of_a_notice():
    zero = style.bar_cells(4000, 0, "0.0")
    assert [item["band"] for item in zero] == [style.TRACK, None] and zero[-1]["text"] == "0.0"
    full = style.bar_cells(4000, 100, "100.0")
    assert [item["band"] for item in full] == [style.BAR, None]
    assert sum(item["width"] for item in zero) == sum(item["width"] for item in full) == 4000
    missing = style.bar_cells(4000, None, "未形成得分")
    assert [(item["band"], item["text"], item["width"]) for item in missing] == [(None, "未形成得分", 4000)]


def test_tables_repeat_header_row_and_have_uniform_header_fill(tmp_path):
    candidate, job = hard_case()
    candidate["parsed"]["fields"] = {"education": {"label": "教育经历", "value": "本科", "status": "人工已确认"}}
    doc = Document(build_report(candidate, job, tmp_path / "tables.docx"))
    for first_header in ("字段", "岗位评价维度", "硬性条件"):
        table = table_headed(doc, first_header)
        properties = table.rows[0]._tr.trPr
        assert properties.find(qn("w:tblHeader")) is not None
        assert {cell_fill(cell) for cell in table.rows[0].cells} == {style.LIGHT}
        assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
        assert table._tbl.tblPr.find(qn("w:tblBorders")) is not None
    # 表格宽度等于版心，不会超出页面。
    section = doc.sections[0]
    room = round((section.page_width - section.left_margin - section.right_margin) / 635)
    for table in doc.tables:
        assert int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) <= room


def test_evidence_is_grouped_by_dimension_with_quote_and_source(tmp_path):
    candidate, job = report_case()
    candidate["score"]["dimensions"][1].update(score=60, reason="提供了行政统筹记录", evidence=[
        {"file": "补充材料.docx", "location": "第2页", "text": "统筹年度会议和后勤保障"},
        {"file": "补充材料.docx", "location": "第2页", "text": "统筹年度会议和后勤保障"}])
    doc = Document(build_report(candidate, job, tmp_path / "evidence.docx"))
    lines = [(p.style.name, p.text, p._p.pPr.find(qn("w:pBdr")) is not None) for p in doc.paragraphs]
    start = lines.index(("Heading 1", "关键证据", True))
    section = lines[start:lines.index(("Heading 1", "人工复核", True))]
    first = section.index(("Heading 3", "沟通协调", False))
    second = section.index(("Heading 3", "行政统筹", False))
    assert section[first + 1][1] == "评分依据：材料说明跨部门协调经历"
    # 引用原文带左边线，下一段是出处（文件 · 位置）。
    assert section[first + 2] == ("Normal", "协调多个部门完成招聘计划", True)
    assert section[first + 3] == ("Normal", "出处：简历.pdf · 第1页", True)
    assert section[second + 2:second + 4] == [("Normal", "统筹年度会议和后勤保障", True),
                                               ("Normal", "出处：补充材料.docx · 第2页", True)]
    # 同一维度内重复的证据只列一次。
    assert [text for _, text, _ in section].count("统筹年度会议和后勤保障") == 1


def test_every_run_sets_the_east_asian_font(tmp_path):
    candidate, job = hard_case()
    doc = Document(build_report(candidate, job, tmp_path / "fonts.docx"))
    parts = [doc.element, doc.sections[0].header._element, doc.sections[0].footer._element]
    runs = [run for part in parts for run in part.iter(qn("w:r"))]
    assert len(runs) > 50
    for run in runs:
        fonts = run.find(qn("w:rPr")).find(qn("w:rFonts"))
        assert fonts.get(qn("w:eastAsia")) == fonts.get(qn("w:ascii")) == "微软雅黑"


def test_reports_keep_the_disclaimer(tmp_path, batch_db):
    candidate, job = report_case()
    assert DISCLAIMER in text_of(Document(build_report(candidate, job, tmp_path / "single.docx")))
    matches = [{"job_id": job["id"], "score": deepcopy(candidate["score"]), "decision": "作为备选岗位"}]
    assert DISCLAIMER in text_of(Document(build_fit_report(candidate, matches, {job["id"]: job}, tmp_path / "fit.docx")))
    add_candidate(batch_db, "C1")
    add_candidate(batch_db, "C2")
    compared = Document(reports.export_comparison_report("T", ["C1", "C2"], "招聘人员", batch_db))
    assert DISCLAIMER in text_of(compared) and "TalentFlow · 人力资源智能体" in compared.sections[0].header.paragraphs[0].text


def test_fit_matrix_is_shaded_by_score_and_marks_the_preferred_job(batch_db):
    high = add_candidate(batch_db, "F1", task="F", decision=None, name="已有人岗记录")
    add_candidate(batch_db, "F2", task="F", decision=None, name="待分析人员")
    low = deepcopy(high)
    low["dimensions"][0]["score"] = 40
    add_match(batch_db, "F1", "J1", low)
    add_match(batch_db, "F1", "J2", high, "推荐至该岗位")
    doc = Document(reports.export_comparison_report("F", ["F1", "F2"], "招聘人员", batch_db))
    matrix = table_headed(doc, "人员")
    first, second = matrix.rows[1].cells[1], matrix.rows[1].cells[2]
    assert first.text == "40.0" and second.text == "★ 85.0\n推荐至该岗位"
    # 得分越高底色越深，都是浅色；没有分析结果的格子不着色。
    assert (cell_fill(first), cell_fill(second)) == (style.heat_fill(40), style.heat_fill(85))
    assert int(cell_fill(first), 16) > int(cell_fill(second), 16) and min(bytes.fromhex(cell_fill(second))) >= 0x90
    assert [cell_fill(cell) for cell in matrix.rows[2].cells[1:]] == [None, None]
    # 优先核验岗位整格加粗；没有人工结论的格子只有得分。
    assert [run.bold for run in second.paragraphs[0].runs] == [True, True]
    assert [run.bold for run in first.paragraphs[0].runs] == [True]


def test_comparison_and_package_tell_unnamed_people_apart(batch_db):
    add_candidate(batch_db, "C1", name="待人工确认")
    add_candidate(batch_db, "C2", name="已确认人员")
    with reports.connect(batch_db) as conn:
        conn.execute("UPDATE candidates SET source_filename=? WHERE id='C1'", ("C:\\private\\应聘材料_甲.pdf",))
    doc = Document(reports.export_comparison_report("T", ["C1", "C2"], "招聘人员", batch_db))
    text = text_of(doc)
    assert "姓名待确认（应聘材料_甲）" in text and "已确认人员" in text
    assert "待人工确认" not in text and "private" not in text
    people = table_headed(doc, "人员及编号")
    assert people.rows[1].cells[0].text == "姓名待确认（应聘材料_甲）\nSAME"
    # 综合匹配度一列带数据条，旁边保留数字。
    assert bands(people.rows[1]) == [style.BAR, style.TRACK] and "85.0" in [cell.text for cell in people.rows[1].cells]
    result = reports.export_batch_reports("T", ["C1", "C2"], "招聘人员", batch_db)
    assert [item["name"] for item in result["exported"]] == ["姓名待确认（应聘材料_甲）", "已确认人员"]
    with ZipFile(result["path"]) as archive:
        documents = sorted(name for name in archive.namelist() if name.endswith(".docx"))
        assert documents[0].startswith("001_姓名待确认（应聘材料_甲）_")
        summary = list(csv.DictReader(io.StringIO(archive.read("审核摘要.csv").decode("utf-8-sig"))))
        assert [row["姓名"] for row in summary] == ["姓名待确认（应聘材料_甲）", "已确认人员"]
        single = text_of(Document(io.BytesIO(archive.read(documents[0]))))
        assert "候选人姓名：姓名待确认（应聘材料_甲）" in single
