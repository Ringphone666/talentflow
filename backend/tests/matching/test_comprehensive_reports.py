from copy import deepcopy

import pytest
from docx import Document

from matching_core.services.report import build_fit_report, build_report, fit_report_readiness, report_readiness


def report_case():
    job = {
        "id": "REPORT-JOB", "name": "综合事务专员", "sequence": "职能类", "level": "5级", "jd_version": "V1",
        "responsibilities": ["负责行政统筹与后勤保障"], "requirements": ["具备沟通协调能力和行政事务经验"],
        "weights": {"沟通协调": 20, "行政统筹": 80}, "hard_conditions": [],
    }
    score = {
        "scoring_method": "ai_semantic_v1", "total": 75,
        "dimensions": [
            {"dimension": "沟通协调", "score": 75, "weight": 20, "weighted": 75,
             "reason": "材料说明跨部门协调经历", "evidence": [{"file": "简历.pdf", "location": "第1页", "text": "协调多个部门完成招聘计划"}]},
            {"dimension": "行政统筹", "score": None, "weight": 80, "weighted": None,
             "reason": "未提供行政制度与后勤工作的具体材料", "evidence": []},
        ],
        "hard_checks": [], "strengths": ["具备跨部门协作经历"], "gaps": [], "risks": [],
        "reasoning_analysis": {"model": "chat-primary", "status": "调用成功", "sections": {
            "strengths": [{"text": "能够协调业务部门推进工作", "evidence": [{"file": "简历.pdf", "location": "第1页", "text": "协调多个部门完成招聘计划"}]}],
            "gaps": [], "risks": [], "questions": [{"text": "请说明协调冲突的处理过程", "evidence": []}],
        }},
    }
    candidate = {"id": "REPORT-CANDIDATE", "full_name": "报告验证人员", "parsed": {"candidate_code": "QA-001", "fields": {}},
                 "score": score, "tags": [], "decision": "保留", "review_comment": "需补充行政统筹项目材料", "parse_status": "人工已复核"}
    return candidate, job


def document_text(path):
    doc = Document(path)
    return "\n".join([*(p.text for p in doc.paragraphs), *(cell.text for table in doc.tables for row in table.rows for cell in row.cells)])


def test_selection_report_uses_whole_job_score_and_preserves_unknowns(tmp_path):
    candidate, job = report_case()
    path = build_report(candidate, job, tmp_path / "selection.docx")
    text = document_text(path)
    assert "综合匹配度：15.0 / 100" in text
    assert "综合匹配度：75.0" not in text
    assert "待补充证据" in text and "未确认贡献" in text
    assert "75.0" in text and "15.0" in text
    for section in ("匹配结论", "目标岗位与评估范围", "关键优势", "主要短板", "风险提示", "建议核验材料", "建议结构化面试问题", "人工最终决策"):
        assert section in text
    assert "简历.pdf" in text and "第1页" in text and "需补充行政统筹项目材料" in text
    assert "归一化" not in text
    assert Document(path).paragraphs[0].style.name == "Title"


def test_fit_report_ranks_consistent_totals_and_omits_stale_confidence(tmp_path):
    candidate, job = report_case()
    second_job = dict(job, id="REPORT-OTHER", name="招聘专员")
    complete = deepcopy(candidate["score"])
    complete["dimensions"][1].update(score=70, weighted=56, reason="提供行政统筹项目记录")
    matches = [
        {"job_id": job["id"], "score": candidate["score"], "decision": "作为备选岗位", "confidence": "高", "recommendation_level": "高度推荐"},
        {"job_id": second_job["id"], "score": complete, "decision": "推荐至该岗位", "review_comment": "已核验项目记录"},
    ]
    path = build_fit_report(candidate, matches, {job["id"]: job, second_job["id"]: second_job}, tmp_path / "fit.docx")
    text = document_text(path)
    assert "当前排序首位岗位：招聘专员" in text
    assert "综合匹配度：71.0 / 100" in text and "综合匹配度：15.0 / 100" in text
    assert "结论可信度" not in text and "高度推荐" not in text
    assert "推荐至该岗位" in text and "已核验项目记录" in text
    assert "建议核验材料" in text and "建议结构化面试问题" in text


@pytest.mark.parametrize("score_value, expected", [(None, "待评估"), (0, "0.0 / 100")])
def test_reports_distinguish_unavailable_from_supported_zero(tmp_path, score_value, expected):
    candidate, job = report_case()
    for item in candidate["score"]["dimensions"]:
        item["score"] = score_value
    text = document_text(build_report(candidate, job, tmp_path / "availability.docx"))
    assert "综合匹配度：" + expected in text
    if score_value is None:
        assert "综合匹配度：0.0" not in text


def test_report_gate_blocks_changed_material_until_reassessment():
    candidate, _ = report_case()
    assert report_readiness(candidate)[0]
    candidate["score"]["requires_reassessment"] = True
    ready, missing = report_readiness(candidate)
    assert not ready and "重新分析" in "".join(missing)
    ready, missing = fit_report_readiness([{"score": candidate["score"], "decision": "推荐至该岗位"}])
    assert not ready and "重新分析" in "".join(missing)


@pytest.mark.parametrize("eligible", [True, False])
def test_fit_report_prioritizes_eligible_job_even_when_top_score_fails_hard_condition(tmp_path, eligible):
    candidate, top_job = report_case()
    top_job["name"] = "高分但条件不足岗位"
    other_job = dict(top_job, id="ELIGIBLE-JOB", name="可优先核验岗位")
    high = deepcopy(candidate["score"])
    lower = deepcopy(candidate["score"])
    for item in high["dimensions"]:
        item.update(score=95, evidence_available=True)
    for item in lower["dimensions"]:
        item.update(score=80, evidence_available=True)
    high["hard_checks"] = [{"condition": "必要资格证书", "status": "不满足", "evidence": "材料明确资格不符"}]
    if not eligible:
        lower["hard_checks"] = [{"condition": "必要资格证书", "status": "无法确认", "evidence": "待补充资格证明"}]
    matches = [{"job_id": top_job["id"], "score": high}, {"job_id": other_job["id"], "score": lower}]
    text = document_text(build_fit_report(candidate, matches, {top_job["id"]: top_job, other_job["id"]: other_job}, tmp_path / "priority.docx"))
    assert "当前排序首位岗位：高分但条件不足岗位" in text
    assert "优先核验岗位：高分但条件不足岗位" not in text
    assert ("优先核验岗位：可优先核验岗位" if eligible else "优先核验岗位：暂无") in text
