"""人岗匹配接口：从人员库取人建立任务、后台分析、人工复核、报告和佐证数据导出，以及各角色的权限。

测试环境没有配置评分模型（MODEL_PROVIDER=local_rules），任务应如实显示“评分未完成”，不得给出分数；
需要评分结果的测试用固定的模型应答代替模型服务，做法与 tests/matching 中的报告回归一致。
"""
import csv
import io
import json
from types import SimpleNamespace
from urllib.parse import unquote
from zipfile import ZipFile

import pytest

from . import support
from .matching.report_fixtures import configure_report_assessment

HR = support.staff("hr_admin", "hr")
ADMIN = support.staff()
BUSINESS = support.staff("business_user", "business")
INTERVIEWER = support.staff("interviewer", "interviewer")
REVIEWER = support.staff("report_reviewer", "report_reviewer")
ZIP = "application/zip"


@pytest.fixture
def library(client, created):
    """两个岗位和一个导入了三份简历的人员批次。"""
    first = support.create_library_job(client, created, "匹配测试岗位甲")
    second = support.create_library_job(client, created, "匹配测试岗位乙", hard_conditions=[],
                                        dimensions=[{"name": "数据分析", "keywords": ["数据", "报表"]}], weights={"数据分析": 100})
    batch = support.create_person_batch(client, created, "匹配测试人员批次")
    files = [(f"匹配人员{index}.txt", support.resume_text(f"匹配人员{index}", f"MATCH-{index}", phone=f"1381234567{index}").encode("utf-8"))
             for index in (1, 2, 3)]
    imported = support.import_persons(client, batch["id"], files)["imported"]
    people = {item["name"]: item["id"] for item in imported}
    return SimpleNamespace(job=first, other_job=second, batch=batch, files=dict(files),
                           people=[people[f"匹配人员{index}"] for index in (1, 2, 3)])


def create_task(client, created, library, *, mode="job_to_people", jobs=None, people=None, name="匹配测试任务", headers=HR):
    body = {"mode": mode, "job_ids": jobs or [library.job["id"]], "person_batch_id": library.batch["id"],
            "person_ids": people or library.people[:2], "name": name}
    response = client.post("/api/matching/tasks", headers=headers, json=body)
    if response.status_code == 200:
        created.tasks.append(response.json()["id"])
    return response


def finished_task(client, created, library, **options):
    response = create_task(client, created, library, **options)
    assert response.status_code == 200, response.text
    return support.wait_task(client, response.json()["id"])


def rows_of(response):
    return list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))


def review(client, candidate_id, headers=HR, **body):
    payload = {"decision": "进入下一轮", "comment": "已核对简历和评分依据", "material_checked": True, **body}
    return client.post(f"/api/matching/candidates/{candidate_id}/review", headers=headers, json=payload)


# ---- 选项与权限 -------------------------------------------------------------------

@pytest.mark.parametrize("role,create,review_,export,delete", [
    ("system_admin", True, True, True, True),
    ("hr_admin", True, True, True, True),
    ("business_user", False, True, True, False),
    ("report_reviewer", False, True, True, False),
    ("interviewer", False, True, False, False),
])
def test_options_follow_the_role(client, role, create, review_, export, delete):
    response = client.get("/api/matching/options", headers=support.staff(role, role))
    assert response.status_code == 200
    options = response.json()
    assert (options["can_create"], options["can_review"], options["can_export"], options["can_delete"]) == (create, review_, export, delete)
    assert options["max_people"] == 20
    assert [item["value"] for item in options["decisions"]] == ["进入下一轮", "保留", "淘汰"]
    assert [item["value"] for item in options["fit_decisions"]] == ["推荐至该岗位", "作为备选岗位", "暂不推荐"]
    assert all(item["label"] for item in options["decisions"] + options["fit_decisions"])


def test_matching_is_closed_to_roles_outside_the_flow(client, created, library):
    task = finished_task(client, created, library)
    candidate = task["people"][0]["candidate_id"]
    requests = [
        ("get", "/api/matching/options", None), ("get", "/api/matching/tasks", None), ("get", f"/api/matching/tasks/{task['id']}", None),
        ("get", f"/api/matching/tasks/{task['id']}/candidates/{candidate}", None),
        ("get", f"/api/matching/tasks/{task['id']}/results", None), ("get", f"/api/matching/tasks/{task['id']}/review-sheet", None),
        ("get", f"/api/matching/tasks/{task['id']}/evidence-package", None),
        ("get", f"/api/matching/tasks/{task['id']}/report?kind=comparison", None),
        ("post", "/api/matching/tasks", {"mode": "job_to_people", "job_ids": [library.job["id"]],
                                         "person_batch_id": library.batch["id"], "person_ids": library.people[:1]}),
        ("post", f"/api/matching/tasks/{task['id']}/retry", None), ("post", f"/api/matching/tasks/{task['id']}/reanalyze", None),
        ("post", f"/api/matching/candidates/{candidate}/review", {"decision": "保留", "comment": "越权复核", "material_checked": True}),
        ("put", f"/api/matching/candidates/{candidate}/material", {"fields": {"education": "越权修改"}}),
        ("delete", f"/api/matching/candidates/{candidate}", None), ("delete", f"/api/matching/tasks/{task['id']}", None),
    ]
    for role in ("question_reviewer", "candidate"):
        for method, path, body in requests:
            response = client.request(method, path, headers=support.staff(role, role), json=body)
            assert response.status_code == 403, (role, method, path)
    for method, path, body in requests:
        assert client.request(method, path, json=body).status_code == 401, (method, path)
    assert support.wait_task(client, task["id"])["people"] == task["people"]


def test_each_role_can_only_do_its_part(client, created, library, monkeypatch):
    configure_report_assessment(monkeypatch)
    task = finished_task(client, created, library)
    first, second = [item["candidate_id"] for item in task["people"]]
    body = {"mode": "job_to_people", "job_ids": [library.job["id"]], "person_batch_id": library.batch["id"],
            "person_ids": library.people[:1]}
    exports = [f"/api/matching/tasks/{task['id']}/results", f"/api/matching/tasks/{task['id']}/review-sheet",
               f"/api/matching/tasks/{task['id']}/evidence-package", f"/api/matching/tasks/{task['id']}/report?kind=comparison"]
    # 面试官、用人部门、报告复核人都不能建立、重试或删除任务。
    for headers in (INTERVIEWER, BUSINESS, REVIEWER):
        assert client.post("/api/matching/tasks", headers=headers, json=body).status_code == 403
        assert client.post(f"/api/matching/tasks/{task['id']}/retry", headers=headers).status_code == 403
        assert client.delete(f"/api/matching/tasks/{task['id']}", headers=headers).status_code == 403
        assert client.delete(f"/api/matching/candidates/{first}", headers=headers).status_code == 403
        assert client.get("/api/matching/tasks", headers=headers).status_code == 200
        assert client.get(f"/api/matching/tasks/{task['id']}", headers=headers).status_code == 200
        assert client.get(f"/api/matching/tasks/{task['id']}/candidates/{first}", headers=headers).status_code == 200
    assert [item["id"] for item in client.get("/api/matching/tasks", headers=HR).json()].count(task["id"]) == 1

    # 面试官参与复核，但不能导出；用人部门和报告复核人可以复核和导出。
    assert review(client, first, headers=INTERVIEWER, decision="保留", comment="面试官已核对").status_code == 200
    for path in exports:
        assert client.get(path, headers=INTERVIEWER).status_code == 403, path
    assert client.get(f"/api/matching/tasks/{task['id']}/report?kind=single&candidate_ids={first}", headers=INTERVIEWER).status_code == 403
    assert review(client, second, headers=BUSINESS, decision="进入下一轮", comment="用人部门已核对").status_code == 200
    for headers in (BUSINESS, REVIEWER):
        for path in exports:
            assert client.get(path, headers=headers).status_code == 200, path
        assert client.get(f"/api/matching/tasks/{task['id']}/report?kind=single&candidate_ids={second}", headers=headers).status_code == 200
    people = {item["candidate_id"]: item for item in client.get(f"/api/matching/tasks/{task['id']}", headers=INTERVIEWER).json()["people"]}
    assert people[first]["decision"] == "保留" and people[first]["reviewer"].endswith("interviewer")
    assert people[second]["decision"] == "进入下一轮" and people[second]["reviewer"].endswith("business")
    assert client.get(f"/api/matching/tasks/{task['id']}/candidates/{first}", headers=HR).json()["reviewer"] == "interviewer"
    sheet = {row["候选人编号"]: row for row in rows_of(client.get(exports[1], headers=BUSINESS))}
    assert sheet["MATCH-1"]["审核人"] == "interviewer" and sheet["MATCH-2"]["审核人"] == "business"
    assert client.post(f"/api/matching/tasks/{task['id']}/reanalyze", headers=BUSINESS).status_code == 200
    support.wait_task(client, task["id"])

    assert client.delete(f"/api/matching/candidates/{second}", headers=HR).json() == {"ok": True}
    assert client.get(f"/api/matching/tasks/{task['id']}/candidates/{second}", headers=HR).status_code == 404
    assert client.delete(f"/api/matching/candidates/{second}", headers=HR).status_code == 404
    assert client.delete(f"/api/matching/tasks/{task['id']}", headers=ADMIN).json() == {"ok": True}
    for path in (f"/api/matching/tasks/{task['id']}", *exports):
        assert client.get(path, headers=HR).status_code == 404, path
    assert client.delete(f"/api/matching/tasks/{task['id']}", headers=HR).status_code == 404
    assert task["id"] not in [item["id"] for item in client.get("/api/matching/tasks", headers=HR).json()]
    # 删除任务不影响人员库里的资料。
    assert client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()["matching_history"] == []
    assert client.get(f"/api/library/person-batches/{library.batch['id']}", headers=HR).json()["parsed_count"] == 3


# ---- 建立任务 --------------------------------------------------------------------

def test_task_needs_valid_jobs_and_parsed_people_of_one_batch(client, created, library):
    other_batch = support.create_person_batch(client, created, "匹配测试另一批次")
    imported = support.import_persons(client, other_batch["id"], [
        ("其他批次人员.txt", support.resume_text("其他批次人员").encode("utf-8")), ("读不出内容.txt", " \n".encode("utf-8"))])
    outsider, unreadable = imported["imported"][0]["id"], imported["failed"][0]["id"]
    archived = support.create_library_job(client, created, "匹配测试停用岗位")
    assert client.post(f"/api/library/jobs/{archived['id']}/archive", headers=HR).status_code == 200
    before = len(client.get("/api/matching/tasks", headers=HR).json())
    cases = [
        (dict(jobs=[library.job["id"], library.other_job["id"]]), 422),          # 岗位选才只选一个岗位
        (dict(mode="people_to_jobs", jobs=[library.job["id"]]), 422),             # 人才适岗至少两个岗位
        (dict(mode="people_to_jobs", jobs=[library.job["id"], library.job["id"]]), 422),
        (dict(jobs=["JOB-NOT-EXISTING"]), 404),
        (dict(jobs=[archived["id"]]), 422),
        (dict(people=[library.people[0], "TLE-NOT-EXISTING"]), 404),
        (dict(people=[library.people[0], outsider]), 422),                        # 不能跨批次混选
        (dict(people=[f"TLE-{index:012d}" for index in range(21)]), 422),         # 超过单个任务的人数上限
    ]
    for options, status in cases:
        assert create_task(client, created, library, **options).status_code == status, options
    body = {"mode": "job_to_people", "job_ids": [library.job["id"]], "person_batch_id": library.batch["id"], "person_ids": library.people[:1]}
    invalid = [dict(body, mode="batch"), dict(body, job_ids=[]), dict(body, person_ids=[]), dict(body, person_batch_id=""),
               dict(body, person_batch_id="TLB-NOT-EXISTING"), dict(body, name="长" * 101),
               dict(body, person_batch_id=other_batch["id"], person_ids=[unreadable]),
               {key: value for key, value in body.items() if key != "person_batch_id"}]
    assert [client.post("/api/matching/tasks", headers=HR, json=item).status_code for item in invalid] == [422, 422, 422, 422, 404, 422, 422, 422]
    assert len(client.get("/api/matching/tasks", headers=HR).json()) == before
    for path in ("/api/matching/tasks/TASK-NOT-EXISTING", "/api/matching/tasks/TASK-NOT-EXISTING/results",
                 "/api/matching/tasks/TASK-NOT-EXISTING/review-sheet", "/api/matching/tasks/TASK-NOT-EXISTING/evidence-package",
                 "/api/matching/tasks/TASK-NOT-EXISTING/candidates/x", "/api/matching/tasks/TASK-NOT-EXISTING/report?kind=package"):
        assert client.get(path, headers=HR).status_code == 404, path
    for action in ("retry", "reanalyze"):
        assert client.post(f"/api/matching/tasks/TASK-NOT-EXISTING/{action}", headers=HR).status_code == 404
    assert client.get("/api/matching/tasks?mode=unknown", headers=HR).status_code == 422


# ---- 没有评分模型时如实呈现 --------------------------------------------------------

def test_without_a_scoring_model_nothing_is_scored(client, created, library):
    mark, core_mark = support.audit_mark(client), support.audit_mark(client, "matching")
    response = create_task(client, created, library, people=library.people[:2] + library.people[:1])
    assert response.status_code == 200
    started = response.json()
    # 任务在后台分析，建立后立即返回；此时是否已经分析完取决于后台进度，不作断言。
    assert started["id"].startswith("TASK-") and started["name"] == "匹配测试任务"
    assert (started["mode"], started["mode_label"]) == ("job_to_people", "岗位选才")
    assert started["jobs"] == [{"id": library.job["id"], "name": "匹配测试岗位甲"}]
    assert started["person_batch_id"] == library.batch["id"] and started["total"] == 2 and started["review"] is None

    task = support.wait_task(client, started["id"])
    assert task["running"] is False and task["status"] == "部分失败"
    assert (task["total"], task["processed"], task["succeeded"], task["failed"]) == (2, 2, 0, 2)
    # 未完成的人员同时给出简历文件名（便于辨认）和人员库编号（便于定位）。
    assert sorted(item["person_id"] for item in task["failures"]) == sorted(library.people[:2])
    assert sorted(item["file"] for item in task["failures"]) == ["匹配人员1.txt", "匹配人员2.txt"]
    assert all(item["stage"] == "analysis" and "AI分析未完成" in item["reason"] for item in task["failures"])
    assert task["review"] == {"passed": 0, "rejected": 0, "kept": 0, "pending": 2, "stale": 0, "exportable": 0, "people": 2}
    assert task["job_standards"] == [{
        "id": library.job["id"], "name": "匹配测试岗位甲", "level": "P3", "sequence": "职能序列", "weight_version": "W-NEW-1.0",
        "dimensions": [{"name": "专业能力", "weight": 60}, {"name": "沟通协调", "weight": 40}], "hard_conditions": ["本科及以上学历"]}]
    assert sorted(item["code"] for item in task["people"]) == ["MATCH-1", "MATCH-2"]
    for person in task["people"]:
        # 不用关键词命中数或其他替代办法凑出分数。
        assert person["total"] is None and person["rank"] is None
        assert person["conclusion"] == "评分未完成" and person["analysis_status"] == "待重试"
        assert person["review_status"] == "待审核" and person["decision"] == "未确认" and person["material_confirmed"] is False
        assert person["report_ready"] is False and any("评分尚未完成" in item for item in person["report_blockers"])
        assert person["strengths"] == [] and person["hard_failed"] == [] and person["hard_unknown"] == ["本科及以上学历"]
        assert "AI分析未完成" in person["note"] and person["person_id"] in library.people[:2]

    person = next(item for item in task["people"] if item["code"] == "MATCH-1")
    detail = client.get(f"/api/matching/tasks/{task['id']}/candidates/{person['candidate_id']}", headers=HR).json()
    assert detail["task"] == {"id": task["id"], "name": "匹配测试任务", "mode": "job_to_people", "running": False}
    assert (detail["code"], detail["name"], detail["person_id"], detail["rank"]) == ("MATCH-1", "匹配人员1", library.people[0], None)
    assert detail["redaction"] == {"status": "已脱敏", "types": ["中国大陆11位手机号"]}
    assert detail["material"]["confirmed"] is False and detail["material"]["material_version"] == 0
    assert detail["report"]["ready"] is False and detail["reports"] == [] and detail["missing_jobs"] == []
    assert "13812345671" not in json.dumps(detail, ensure_ascii=False)
    assert len(detail["assessments"]) == 1
    assessment = detail["assessments"][0]
    assert assessment["job_id"] == library.job["id"] and assessment["decision"] is None
    assert assessment["total"] is None and assessment["coverage_percent"] == 0 and assessment["recommended"] is False
    assert assessment["conclusion"] == "评分未完成" and assessment["analysis_status"] == "待重试" and assessment["strengths"] == []
    assert [(item["dimension"], item["weight"], item["score"], item["weighted"], item["status"]) for item in assessment["dimensions"]] == [
        ("专业能力", 60, None, None, "待核实"), ("沟通协调", 40, None, None, "待核实")]
    assert [(item["condition"], item["status"]) for item in assessment["hard_checks"]] == [("本科及以上学历", "无法确认")]
    assert assessment["model"]["status"] == "未成功" and assessment["model"]["reason"]
    assert any("AI评分未完成" in item for item in assessment["risks"])

    listed = client.get("/api/matching/tasks?mode=job_to_people", headers=HR).json()
    assert next(item for item in listed if item["id"] == task["id"])["review"] == task["review"]
    assert task["id"] not in [item["id"] for item in client.get("/api/matching/tasks?mode=people_to_jobs", headers=HR).json()]
    assert client.get(f"/api/library/jobs/{library.job['id']}", headers=HR).json()["matching_task_count"] == 1
    history = client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()["matching_history"]
    assert [(item["task_id"], item["mode"], item["job_name"], item["total"], item["conclusion"], item["decision"]) for item in history] == [
        (task["id"], "job_to_people", "匹配测试岗位甲", None, "评分未完成", None)]
    overview = client.get("/api/overview", headers=HR).json()
    assert overview["matching"]["with_failures"] >= 1 and overview["matching"]["pending_review"] >= 2
    assert any("未完成" in item["title"] for item in overview["todo"])

    # 重试和重新分析不会凭空产生分数。
    for action in ("retry", "reanalyze"):
        assert client.post(f"/api/matching/tasks/{task['id']}/{action}", headers=HR).json() == {"ok": True}
        again = support.wait_task(client, task["id"])
        assert again["status"] == "部分失败" and [item["total"] for item in again["people"]] == [None, None]
    platform = [item["action"] for item in reversed(support.audits_since(client, mark)) if item["target_id"] == task["id"]]
    assert platform == ["matching_task_created", "matching_task_retried", "matching_task_reanalyzed"]
    core = [item for item in support.audits_since(client, core_mark, source="matching") if item["target_id"].startswith(task["id"])]
    assert {"人员库选用", "模型匹配分析", "人员材料复用"} <= {item["action"] for item in core}
    assert {item["actor"] for item in core} == {"hr"} and {item["role"] for item in core} == {"招聘人员"}


def test_exports_state_that_scoring_is_not_finished(client, created, library):
    task = finished_task(client, created, library)
    first, second = sorted(task["people"], key=lambda item: item["code"])
    path = f"/api/matching/tasks/{task['id']}"

    results = client.get(path + "/results", headers=HR)
    assert results.status_code == 200
    body = results.json()
    assert body["task"]["id"] == task["id"] and body["task"]["mode"] == "job_to_people" and body["task"]["status"] == "部分失败"
    assert body["notice"] and body["generated_at"]
    assert [(job["id"], job["name"], job["weights"]) for job in body["jobs"]] == [(library.job["id"], "匹配测试岗位甲", {"专业能力": 60, "沟通协调": 40})]
    assert sorted(item["candidate_code"] for item in body["people"]) == ["MATCH-1", "MATCH-2"]
    for person in body["people"]:
        assert person["rank"] is None and person["report_ready"] is False and person["review_status"] == "待审核"
        assert [(item["total"], item["conclusion"], item["decision"]) for item in person["assessments"]] == [(None, "评分未完成", None)]
        assert all(dimension["score"] is None for item in person["assessments"] for dimension in item["dimensions"])
        assert person["material"]["source_filename"] in library.files
    serialized = json.dumps(body, ensure_ascii=False)
    # 对外输出不含路径和未脱敏的联系方式。
    assert "original_path" not in serialized and "redacted_path" not in serialized and "test-data" not in serialized
    assert "13812345671" not in serialized
    assert "13812345672" not in serialized

    sheet = client.get(path + "/review-sheet", headers=HR)
    assert sheet.status_code == 200 and sheet.headers["content-type"].startswith("text/csv")
    assert unquote(sheet.headers["x-download-name"]) == f"{task['id']}_审核明细.csv"
    rows = {row["候选人编号"]: row for row in rows_of(sheet)}
    assert set(rows) == {"MATCH-1", "MATCH-2"} and "candidate_id" not in rows["MATCH-1"]
    assert list(rows["MATCH-1"]) == ["候选人编号", "姓名", "原始文件", "综合匹配度", "匹配结论", "分析状态", "审核状态", "人工结论",
                                    "资料确认", "已审岗位", "复核意见", "审核人", "审核时间", "报告状态", "处理说明"]
    assert (rows["MATCH-1"]["姓名"], rows["MATCH-1"]["综合匹配度"], rows["MATCH-1"]["匹配结论"]) == ("匹配人员1", "", "评分未完成")
    assert rows["MATCH-1"]["分析状态"] == "待重试" and rows["MATCH-1"]["人工结论"] == "未确认" and "评分尚未完成" in rows["MATCH-1"]["报告状态"]
    only = rows_of(client.get(path + f"/review-sheet?candidate_ids={second['candidate_id']}", headers=HR))
    assert [row["候选人编号"] for row in only] == ["MATCH-2"]

    # 评分没有完成，正式报告不能导出；多人对比如实标注待分析。
    single = client.get(path + f"/report?kind=single&candidate_ids={first['candidate_id']}", headers=HR)
    assert single.status_code == 422 and "评分尚未完成" in single.json()["detail"]
    assert client.get(path + "/report?kind=single", headers=HR).status_code == 422
    package = client.get(path + "/report?kind=package", headers=HR)
    assert package.status_code == 409 and "匹配人员1" in package.json()["detail"]
    assert client.get(path + f"/report?kind=comparison&candidate_ids={first['candidate_id']}", headers=HR).status_code == 422
    assert client.get(path + "/report?kind=summary", headers=HR).status_code == 422
    assert client.get(path + "/report?kind=single&candidate_ids=NOT-IN-TASK", headers=HR).status_code == 422
    comparison = client.get(path + "/report?kind=comparison", headers=HR)
    assert comparison.status_code == 200 and comparison.headers["content-type"] == support.DOCX
    from docx import Document
    document = Document(io.BytesIO(comparison.content))
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert "匹配人员1" in text and "匹配人员2" in text and "待分析" in text and "未确认" in text
    assert "审核中" not in text and "已分析" not in text
    assert "80.0" not in text and "高度匹配" not in text
    assert [item["file"] for item in client.get(path + f"/candidates/{first['candidate_id']}", headers=HR).json()["reports"]] == [
        unquote(comparison.headers["x-download-name"])]


def test_evidence_package_contains_the_whole_trail(client, created, library):
    task = finished_task(client, created, library)
    first, second = sorted(task["people"], key=lambda item: item["code"])
    path = f"/api/matching/tasks/{task['id']}/evidence-package"
    core_mark = support.audit_mark(client, "matching")
    response = client.get(path, headers=HR)
    assert response.status_code == 200 and response.headers["content-type"] == ZIP
    name = unquote(response.headers["x-download-name"])
    assert name.startswith("匹配测试任务_佐证数据_") and name.endswith(".zip")
    with ZipFile(io.BytesIO(response.content)) as archive:
        names = archive.namelist()
        assert names == ["说明.json", "结果.json", "人员/MATCH-1_匹配人员1/脱敏材料.txt", "人员/MATCH-1_匹配人员1/结构化资料.json",
                         "人员/MATCH-2_匹配人员2/脱敏材料.txt", "人员/MATCH-2_匹配人员2/结构化资料.json", "操作记录.csv"]
        manifest = json.loads(archive.read("说明.json"))
        results = json.loads(archive.read("结果.json"))
        redacted = archive.read("人员/MATCH-1_匹配人员1/脱敏材料.txt").decode("utf-8")
        structured = json.loads(archive.read("人员/MATCH-1_匹配人员1/结构化资料.json"))
        records = list(csv.DictReader(io.StringIO(archive.read("操作记录.csv").decode("utf-8-sig"))))
        everything = b"".join(archive.read(item) for item in names).decode("utf-8", errors="ignore")
    assert manifest["task"]["id"] == task["id"] and manifest["people"] == 2 and manifest["include_originals"] is False
    assert manifest["generated_by"] == "hr" and manifest["notice"] and len(manifest["contents"]) == 6
    assert results["task"]["id"] == task["id"] and len(results["people"]) == 2
    assert all(item["assessments"][0]["total"] is None for item in results["people"])
    assert "三年招聘与人员调配工作" in redacted and "138****5671" in redacted
    assert structured["candidate_code"] == "MATCH-1" and structured["fields"]["education"]["value"] == "本科，人力资源管理专业"
    assert any(item["text"] == "姓名：匹配人员1" for item in structured["raw_evidence"])
    # 默认不含原始简历，包内没有未脱敏的联系方式和本机路径。
    assert "13812345671" not in everything and "13812345672" not in everything and "test-data" not in everything
    assert list(records[0]) == ["操作时间", "操作类型", "对象类型", "对象", "操作人", "变更前", "变更后"]
    assert [item["操作类型"] for item in records][0] == "人员库选用"
    assert {"模型匹配分析", "人员材料复用"} <= {item["操作类型"] for item in records}
    assert all(item["对象"].startswith(task["id"]) for item in records) and {item["操作人"] for item in records} == {"招聘人员｜hr"}
    exported = [item for item in support.audits_since(client, core_mark, source="matching") if item["action"] == "佐证数据导出"]
    assert len(exported) == 1 and exported[0]["target_id"] == task["id"] and exported[0]["actor"] == "hr"
    assert exported[0]["detail"]

    # 指定人员并明确要求时才打包原始简历。
    selected = client.get(path, headers=HR, params={"candidate_ids": second["candidate_id"], "include_originals": "true"})
    assert selected.status_code == 200
    with ZipFile(io.BytesIO(selected.content)) as archive:
        assert archive.namelist() == ["说明.json", "结果.json", "人员/MATCH-2_匹配人员2/脱敏材料.txt", "人员/MATCH-2_匹配人员2/结构化资料.json",
                                      "人员/MATCH-2_匹配人员2/原始材料/匹配人员2.txt", "操作记录.csv"]
        assert archive.read("人员/MATCH-2_匹配人员2/原始材料/匹配人员2.txt") == library.files["匹配人员2.txt"]
        manifest = json.loads(archive.read("说明.json"))
        assert manifest["people"] == 1 and manifest["include_originals"] is True
        assert "佐证数据导出" in {row["操作类型"] for row in csv.DictReader(io.StringIO(archive.read("操作记录.csv").decode("utf-8-sig")))}
    assert client.get(path, headers=HR, params={"candidate_ids": "NOT-IN-TASK"}).status_code == 422
    assert first["candidate_id"] != second["candidate_id"]


# ---- 有评分结果时的复核与报告 -------------------------------------------------------

def test_scored_people_are_ranked_reviewed_and_reported(client, created, library, monkeypatch):
    configure_report_assessment(monkeypatch)
    task = finished_task(client, created, library, people=library.people)
    assert task["status"] == "已完成" and (task["total"], task["succeeded"], task["failed"], task["failures"]) == (3, 3, 0, [])
    assert task["review"] == {"passed": 0, "rejected": 0, "kept": 0, "pending": 3, "stale": 0, "exportable": 0, "people": 3}
    assert sorted(item["rank"] for item in task["people"]) == [1, 2, 3]
    for person in task["people"]:
        assert person["total"] == 80 and person["analysis_status"] == "分析完成" and person["conclusion"] == "硬性条件待核实"
        assert person["report_ready"] is False and person["note"] == "" and len(person["strengths"]) == 2
    people = {item["code"]: item["candidate_id"] for item in task["people"]}
    first, second, third = people["MATCH-1"], people["MATCH-2"], people["MATCH-3"]
    path = f"/api/matching/tasks/{task['id']}"
    detail = client.get(f"{path}/candidates/{first}", headers=HR).json()
    assessment = detail["assessments"][0]
    assert assessment["total"] == 80 and assessment["coverage_percent"] == 100 and assessment["model"]["status"] == "调用成功"
    assert [(item["dimension"], item["score"], item["weighted"], item["status"]) for item in assessment["dimensions"]] == [
        ("专业能力", 80, 48, "已评估"), ("沟通协调", 80, 32, "已评估")]
    assert all(item["reason"] and item["evidence"] for item in assessment["dimensions"])
    # 硬性条件没有核实前不作为推荐人选。
    assert assessment["recommended"] is False and [item["status"] for item in assessment["hard_checks"]] == ["无法确认"]

    # 复核结论、意见和“已核对资料”缺一不可。
    invalid = [dict(material_checked=False), dict(decision="录用"), dict(decision="推荐至该岗位"), dict(decision=""), dict(comment="好"),
               dict(comment=""), dict(expected_material_version=5)]
    for body in invalid:
        assert review(client, first, **body).status_code == 422, body
    incomplete = client.post(f"/api/matching/candidates/{first}/review", headers=HR, json={"decision": "保留", "comment": "缺少核对确认"})
    assert incomplete.status_code == 422
    assert review(client, "TASK-NOT-EXISTING-X").status_code == 404
    unchanged = client.get(f"{path}/candidates/{first}", headers=HR).json()
    assert unchanged["assessments"][0]["decision"] is None and unchanged["review_status"] == "待审核"
    assert unchanged["material"]["confirmed"] is False
    assert client.get(f"{path}/report?kind=single&candidate_ids={first}", headers=HR).status_code == 422
    assert client.get(f"{path}/report?kind=package", headers=HR).status_code == 409

    mark = support.audit_mark(client)
    assert review(client, first, expected_material_version=0).json() == {"ok": True}
    assert review(client, second, decision="保留", comment="学历证明待补充").json() == {"ok": True}
    assert review(client, third, decision="淘汰", comment="与岗位要求差距较大").json() == {"ok": True}
    saved = client.get(f"{path}/candidates/{first}", headers=HR).json()
    assert saved["review_status"] == "通过" and saved["reviewer"] == "hr" and saved["reviewed_at"]
    assert saved["assessments"][0]["decision"] == "进入下一轮" and saved["assessments"][0]["review_comment"] == "已核对简历和评分依据"
    # 提交结论时一并确认资料。
    assert saved["material"]["confirmed"] is True and saved["material"]["material_version"] == 1
    assert saved["report"] == {"ready": True, "blockers": []}
    reviewed = client.get(path, headers=HR).json()
    assert reviewed["review"] == {"passed": 1, "rejected": 1, "kept": 1, "pending": 0, "stale": 0, "exportable": 3, "people": 3}
    assert {item["code"]: item["review_status"] for item in reviewed["people"]} == {"MATCH-1": "通过", "MATCH-2": "保留", "MATCH-3": "未通过"}
    assert [(item["action"], item["detail"]) for item in reversed(support.audits_since(client, mark))] == [
        ("matching_review_saved", "进入下一轮"), ("matching_review_saved", "保留"), ("matching_review_saved", "淘汰")]
    # 人员库里的资料不因任务内的确认而改变。
    assert client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()["confirmed"] is False
    history = client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()["matching_history"]
    assert [(item["task_id"], item["total"], item["review_status"], item["decision"]) for item in history] == [
        (task["id"], 80, "通过", "进入下一轮")]

    from docx import Document

    def text_of(content):
        document = Document(io.BytesIO(content))
        return "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])

    single = client.get(f"{path}/report?kind=single&candidate_ids={first}", headers=HR)
    assert single.status_code == 200 and single.headers["content-type"] == support.DOCX
    assert unquote(single.headers["x-download-name"]).startswith("人岗匹配报告_匹配人员1_")
    report = text_of(single.content)
    assert "匹配人员1" in report and "匹配测试岗位甲" in report and "进入下一轮" in report and "已核对简历和评分依据" in report
    assert "13812345671" not in report
    assert client.get(f"{path}/report?kind=single&candidate_ids={first},{second}", headers=HR).status_code == 422
    package = client.get(f"{path}/report?kind=package&candidate_ids={first},{third}", headers=BUSINESS)
    assert package.status_code == 200 and package.headers["content-type"] == ZIP
    with ZipFile(io.BytesIO(package.content)) as archive:
        documents = [item for item in archive.namelist() if item.endswith(".docx")]
        assert len(documents) == 2 and sorted(archive.namelist()) == sorted(documents + ["审核摘要.csv"])
        summary = list(csv.DictReader(io.StringIO(archive.read("审核摘要.csv").decode("utf-8-sig"))))
        assert [(row["候选人编号"], row["人工结论"], row["导出结果"]) for row in summary] == [
            ("MATCH-1", "进入下一轮", "已导出"), ("MATCH-3", "淘汰", "已导出")]
        assert "匹配人员3" in text_of(archive.read(documents[1]))
    comparison = client.get(f"{path}/report?kind=comparison", headers=HR)
    assert comparison.status_code == 200
    compared = text_of(comparison.content)
    assert all(value in compared for value in ("匹配人员1", "匹配人员2", "匹配人员3", "80.0", "进入下一轮", "保留", "淘汰"))

    sheet = {row["候选人编号"]: row for row in rows_of(client.get(f"{path}/review-sheet", headers=HR))}
    assert (sheet["MATCH-1"]["综合匹配度"], sheet["MATCH-1"]["人工结论"], sheet["MATCH-1"]["审核人"]) == ("80.0", "进入下一轮", "hr")
    assert sheet["MATCH-1"]["报告状态"] == "可导出" and sheet["MATCH-2"]["复核意见"] == "学历证明待补充"
    results = {item["candidate_code"]: item for item in client.get(f"{path}/results", headers=HR).json()["people"]}
    assert results["MATCH-1"]["assessments"][0]["total"] == 80 and results["MATCH-1"]["assessments"][0]["decision"] == "进入下一轮"
    assert results["MATCH-1"]["report_ready"] is True and results["MATCH-1"]["reviewer"] == "hr" and results["MATCH-3"]["review_status"] == "未通过"
    evidence = client.get(f"{path}/evidence-package", headers=HR)
    with ZipFile(io.BytesIO(evidence.content)) as archive:
        # 单人报告放在本人的文件夹里；报告包和多人对比报告涉及多人，只在“报告/”里放一份。
        names = archive.namelist()
        assert [item for item in names if item.startswith("人员/MATCH-1_匹配人员1/报告/")] == [
            f"人员/MATCH-1_匹配人员1/报告/{unquote(single.headers['x-download-name'])}"]
        assert sorted(item for item in names if item.startswith("报告/")) == sorted(
            f"报告/{unquote(item.headers['x-download-name'])}" for item in (package, comparison))
        assert not [item for item in names if item.startswith("人员/MATCH-3_") and "多人对比" in item]
        assert archive.read(f"人员/MATCH-1_匹配人员1/报告/{unquote(single.headers['x-download-name'])}") == single.content
        assert "人工决策" in archive.read("操作记录.csv").decode("utf-8-sig")
    assert len(client.get(f"{path}/candidates/{first}", headers=HR).json()["reports"]) == 3


def test_material_change_requires_reanalysis_before_review(client, created, library, monkeypatch):
    configure_report_assessment(monkeypatch)
    task = finished_task(client, created, library, people=library.people[:1])
    candidate = task["people"][0]["candidate_id"]
    path = f"/api/matching/tasks/{task['id']}"
    assert review(client, candidate).status_code == 200
    rejected = [{"fields": {"salary": "面议"}}, {"fields": {"education": "硕士"}, "expected_material_version": 0},
                {"fields": {"additional_material": "手机号 13912345678"}}, {"fields": {}, "full_name": " "}]
    for body in rejected:
        assert client.put(f"/api/matching/candidates/{candidate}/material", headers=HR, json=body).status_code == 422, body
    assert client.put("/api/matching/candidates/NOT-EXISTING/material", headers=HR, json={"fields": {}}).status_code == 404
    assert client.get(path, headers=HR).json()["people"][0]["review_status"] == "通过"

    mark = support.audit_mark(client)
    changed = client.put(f"/api/matching/candidates/{candidate}/material", headers=BUSINESS, json={
        "fields": {"education": "硕士，人力资源管理专业"}, "full_name": "匹配人员一", "expected_material_version": 1})
    assert changed.status_code == 200
    body = changed.json()
    # 资料有改动：原来的评分和人工结论失效，需要重新分析后再复核。
    assert body["requires_reanalysis"] is True and body["warning"] == ""
    assert body["material"]["full_name"] == "匹配人员一" and body["material"]["material_version"] == 2
    assert body["material"]["confirmed"] is False and body["material"]["updated_at"]
    fields = {item["key"]: item for item in body["material"]["fields"]}
    assert fields["education"]["value"] == "硕士，人力资源管理专业" and fields["education"]["status"] == "人工已修改，待确认"
    assert [item["action"] for item in support.audits_since(client, mark)] == ["matching_material_updated"]
    person = client.get(path, headers=HR).json()["people"][0]
    assert person["name"] == "匹配人员一" and person["decision"] == "未确认" and person["review_status"] == "需重新审核"
    assert person["analysis_status"] == "资料已更新，待重分析" and person["conclusion"] == "资料已变更，待重新分析"
    assert person["report_ready"] is False and person["total"] == 80
    detail = client.get(f"{path}/candidates/{candidate}", headers=HR).json()
    assert detail["assessments"][0]["stale"] is True and detail["assessments"][0]["stale_reason"]
    assert client.get(path, headers=HR).json()["review"]["stale"] == 1
    assert review(client, candidate, decision="保留").status_code == 422
    assert client.get(f"{path}/report?kind=single&candidate_ids={candidate}", headers=HR).status_code == 422
    assert client.get("/api/overview", headers=HR).json()["matching"]["stale"] >= 1
    # 任务内的修改不回写人员库。
    library_person = client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()
    assert library_person["name"] == "匹配人员1" and library_person["education"] == "本科，人力资源管理专业"

    assert client.post(f"{path}/reanalyze", headers=REVIEWER).json() == {"ok": True}
    person = support.wait_task(client, task["id"])["people"][0]
    assert person["analysis_status"] == "分析完成" and person["review_status"] == "待审核" and person["total"] == 80
    assert review(client, candidate, decision="保留", comment="按修正后的资料复核").status_code == 200
    person = client.get(path, headers=HR).json()["people"][0]
    assert (person["decision"], person["review_status"], person["material_confirmed"], person["report_ready"]) == ("保留", "保留", True, True)
    exported = client.get(f"{path}/report?kind=single&candidate_ids={candidate}", headers=HR)
    assert exported.status_code == 200
    from docx import Document
    document = Document(io.BytesIO(exported.content))
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert "匹配人员一" in text and "硕士，人力资源管理专业" in text and "按修正后的资料复核" in text


def test_rejected_review_changes_nothing(client, created, library, monkeypatch):
    configure_report_assessment(monkeypatch)
    task = finished_task(client, created, library, people=library.people[:1])
    candidate = task["people"][0]["candidate_id"]
    assert client.put(f"/api/matching/candidates/{candidate}/material", headers=HR,
                      json={"fields": {"education": "硕士，人力资源管理专业"}}).status_code == 200
    path = f"/api/matching/tasks/{task['id']}/candidates/{candidate}"
    before = client.get(path, headers=HR).json()["material"]
    assert before["confirmed"] is False
    assert review(client, candidate, decision="保留").status_code == 422
    after = client.get(path, headers=HR).json()["material"]
    assert (after["confirmed"], after["material_version"]) == (False, before["material_version"])


def test_people_to_jobs_compares_one_person_across_jobs(client, created, library, monkeypatch):
    configure_report_assessment(monkeypatch)
    jobs = [library.job["id"], library.other_job["id"]]
    response = create_task(client, created, library, mode="people_to_jobs", jobs=jobs, people=library.people[:1], name="")
    assert response.status_code == 200
    started = response.json()
    # 没有填写任务名称时沿用人员批次的名称。
    assert started["id"].startswith("FIT-") and started["name"] == "匹配测试人员批次"
    assert (started["mode"], started["mode_label"]) == ("people_to_jobs", "人才适岗")
    assert [item["id"] for item in started["jobs"]] == jobs
    task = support.wait_task(client, started["id"])
    assert task["status"] == "已完成" and [item["id"] for item in task["job_standards"]] == jobs
    person = task["people"][0]
    candidate = person["candidate_id"]
    assert person["reviewed_jobs"] == "0/2" and person["review_status"] == "待审核" and person["total"] == 80
    assert person["matrix"] == {
        jobs[0]: {"total": 80, "decision": None, "analysis_status": "分析完成", "conclusion": "硬性条件待核实"},
        jobs[1]: {"total": 80, "decision": None, "analysis_status": "分析完成", "conclusion": "高度匹配"}}
    # 分数相同的两个岗位里，硬性条件已无疑问的岗位才是推荐岗位。
    assert person["preferred_job_id"] == jobs[1]
    path = f"/api/matching/tasks/{task['id']}"
    assert task["id"] in [item["id"] for item in client.get("/api/matching/tasks?mode=people_to_jobs", headers=HR).json()]
    assert task["id"] not in [item["id"] for item in client.get("/api/matching/tasks?mode=job_to_people", headers=HR).json()]

    invalid = [dict(decision="推荐至该岗位"), dict(decision="进入下一轮", job_id=jobs[0]), dict(decision="推荐至该岗位", job_id="JOB-NOT-IN-TASK")]
    assert [review(client, candidate, **body).status_code for body in invalid] == [422, 422, 404]
    assert client.get(f"{path}/report?kind=single&candidate_ids={candidate}", headers=HR).status_code == 422
    assert review(client, candidate, decision="推荐至该岗位", job_id=jobs[1], comment="数据分析经历与岗位相符").status_code == 200
    assert review(client, candidate, decision="暂不推荐", job_id=jobs[0], comment="学历证明尚未核实").status_code == 200
    detail = client.get(f"{path}/candidates/{candidate}", headers=HR).json()
    assert detail["task"]["mode"] == "people_to_jobs" and detail["review_status"] == "通过" and detail["missing_jobs"] == []
    assert {item["job_id"]: (item["job_name"], item["decision"], item["total"]) for item in detail["assessments"]} == {
        jobs[0]: ("匹配测试岗位甲", "暂不推荐", 80), jobs[1]: ("匹配测试岗位乙", "推荐至该岗位", 80)}
    assert detail["report"] == {"ready": True, "blockers": []} and detail["material"]["confirmed"] is True
    person = client.get(path, headers=HR).json()["people"][0]
    assert person["reviewed_jobs"] == "2/2" and person["matrix"][jobs[1]]["decision"] == "推荐至该岗位"
    assert "匹配测试岗位乙：推荐至该岗位" in person["decision"] and "匹配测试岗位甲：暂不推荐" in person["decision"]
    history = client.get(f"/api/library/persons/{library.people[0]}", headers=HR).json()["matching_history"]
    assert [(item["task_id"], item["mode"]) for item in history] == [(task["id"], "people_to_jobs")]
    assert client.get(f"/api/library/jobs/{jobs[1]}", headers=HR).json()["matching_task_count"] == 1

    exported = client.get(f"{path}/report?kind=single&candidate_ids={candidate}", headers=HR)
    assert exported.status_code == 200 and unquote(exported.headers["x-download-name"]).startswith("人才适岗推荐报告_匹配人员1_")
    from docx import Document
    document = Document(io.BytesIO(exported.content))
    text = "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])
    assert "匹配测试岗位甲" in text and "匹配测试岗位乙" in text and "数据分析经历与岗位相符" in text
    results = client.get(f"{path}/results", headers=HR).json()
    assert results["task"]["mode"] == "people_to_jobs" and [job["id"] for job in results["jobs"]] == jobs
    assert sorted(item["job_id"] for item in results["people"][0]["assessments"]) == sorted(jobs)
    # 岗位标准改动后，已有结果标记为需要重新分析。
    updated = client.put(f"/api/library/jobs/{jobs[1]}", headers=HR, json=support.job_payload(
        "匹配测试岗位乙", hard_conditions=[], dimensions=[{"name": "数据分析", "keywords": ["数据", "报表", "指标"]}], weights={"数据分析": 100}))
    assert updated.status_code == 200 and updated.json()["scorecard_change"]["changed"] is True
    stale = client.get(path, headers=HR).json()
    assert stale["review"]["stale"] == 1 and stale["people"][0]["review_status"] == "需重新审核"
    assert stale["people"][0]["report_ready"] is False


# ---- 与AI面试的衔接 ----------------------------------------------------------------

def test_matching_result_is_shown_as_hard_indicators_in_interviews(client, created, library, monkeypatch):
    interviewer = support.staff("interviewer", "interviewer")
    batch = support.create_batch(client, created, "MATCH-LINK", "匹配衔接测试批次", library_job_id=library.job["id"], mode="listening_only")
    linked = support.create_candidate(client, created, batch["id"], "匹配人员1", library_person_id=library.people[0])
    plain = support.create_candidate(client, created, batch["id"], "没有档案的候选人")

    def indicators(candidate):
        return client.get(f"/api/listening/{candidate['id']}", headers=interviewer).json()["hard_indicators"]

    missing = indicators(plain)
    assert missing["available"] is False and missing["reason"] and missing["profile"] is None
    assert missing["facts"] == [] and missing["hard_checks"] == [] and missing["matching"] is None

    # 还没有做过人岗匹配：只按岗位硬性条件做关键词初核，并标明需要人工核实。
    before = indicators(linked)
    assert before["available"] is True and before["reason"] == "" and before["counted_in_score"] is False and before["notice"]
    assert before["profile"]["person_id"] == library.people[0] and before["profile"]["confirmed"] is False
    assert [item["key"] for item in before["facts"]] == ["education", "work_experience", "tenure", "skills"]
    assert {item["key"]: item["value"] for item in before["facts"]}["work_experience"] == "三年招聘与人员调配工作"
    assert before["matching"] is None
    assert [(item["condition"], item["source"]) for item in before["hard_checks"]] == [("本科及以上学历", "关键词初核，待人工核实")]
    assert before["hard_checks"][0]["status"] in ("满足", "不满足", "无法确认")

    configure_report_assessment(monkeypatch)
    task = finished_task(client, created, library, people=library.people[:1])
    after = indicators(linked)
    assert after["available"] is True and after["counted_in_score"] is False and after["facts"] == before["facts"]
    assert after["matching"] == {"task_id": task["id"], "candidate_id": task["people"][0]["candidate_id"], "total": 80,
                                 "conclusion": "硬性条件待核实", "coverage_percent": 100}
    assert [(item["condition"], item["status"], item["source"]) for item in after["hard_checks"]] == [
        ("本科及以上学历", "无法确认", "人岗匹配分析")]
    assert after["hard_checks"][0]["evidence"]
    # 其他岗位的匹配结果不算在本岗位名下。
    other_batch = support.create_batch(client, created, "MATCH-LINK-2", "另一岗位的批次", library_job_id=library.other_job["id"],
                                       mode="listening_only")
    elsewhere = support.create_candidate(client, created, other_batch["id"], "匹配人员1", library_person_id=library.people[0])
    assert indicators(elsewhere)["matching"] is None and indicators(elsewhere)["hard_checks"] == []

    # 岗位标准改动后，原来的匹配结果不再作为硬指标的依据。
    assert client.put(f"/api/library/jobs/{library.job['id']}", headers=HR, json=support.job_payload(
        "匹配测试岗位甲", requirements=["本科及以上学历", "五年以上人力资源工作经验"])).status_code == 200
    stale = indicators(linked)
    assert stale["matching"] is None
    assert [(item["condition"], item["source"]) for item in stale["hard_checks"]] == [("本科及以上学历", "关键词初核，待人工核实")]
