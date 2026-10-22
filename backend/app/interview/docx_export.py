"""AI 面试报告、综合报告和候选人对比的 Word 导出。

版式（页面、页眉页脚、摘要卡、表格、数据条）与人岗匹配报告共用 matching_core 里的 report_style；
报告内容全部来自已保存的报告数据，这里不做任何评分或改写。
"""
import io
import math

from docx.shared import Twips

from matching_core.services import report_style as style

NOTICE = "AI辅助生成，须人工复核。本文档用于辅助人工评估，不构成录用或淘汰决定。"
NUMERALS = ["", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]
SPEAKERS = {"interviewer": "面试官", "candidate": "候选人", "dialogue": "混合对话待核对"}
# 连续对话里按对话语境逐段区分出的发言人。
ROLES = {"interviewer": "面试官", "candidate": "候选人", "unknown": "未能判断"}
PRESERVED = "证据保全降级"
FULL_SCORE = 5


def _number(value):
    """报告内容可能经过人工修改，只有数值才画数据条和雷达图。"""
    return float(value) if type(value) in (int, float) and math.isfinite(value) else None


def _percent(value) -> str:
    return "—" if value is None else f"{value:g}%"


def _section(doc, number, title):
    return style.heading(doc, f"{NUMERALS[number] if number < len(NUMERALS) else number}、{title}")


def _transcript(doc, label, text):
    """问答原文：左边线标出，与分析文字区分。"""
    return style.paragraph(doc, [(label, 10, True, style.INK), (str(text), 10, False, style.INK)], after=4, line=1.3,
                           left=420, borders={"left": (style.BRAND, 12, 9)})


# ---- 六维雷达图 ------------------------------------------------------------------

def radar_png(labels, values, maximum=FULL_SCORE):
    """六维雷达图（PNG）。没有得分的维度不连线，也不按 0 分画。

    有得分的维度不足三个，或绘制失败时返回 None，报告里不放图。
    """
    known = [index for index, value in enumerate(values) if value is not None]
    if len(labels) < 3 or len(known) < 3:
        return None
    names = [str(label) if len(str(label)) <= 12 else str(label)[:11] + "…" for label in labels]
    try:
        import fitz
        with fitz.open() as document:
            return _draw_radar(fitz, document.new_page(width=440, height=284), names, values, known, maximum)
    except Exception:
        return None


def _draw_radar(fitz, page, labels, values, known, maximum):
    brand, grid = (0x19 / 255, 0x74 / 255, 0x65 / 255), (0.79, 0.86, 0.85)
    ink, muted = (0.12, 0.16, 0.15), (0.36, 0.40, 0.45)
    count, radius, center = len(labels), 96, fitz.Point(220, 144)

    def angle(index):
        return -math.pi / 2 + 2 * math.pi * index / count

    def point(index, ratio):
        return fitz.Point(center.x + radius * ratio * math.cos(angle(index)),
                          center.y + radius * ratio * math.sin(angle(index)))

    shape = page.new_shape()
    for level in range(maximum, 0, -1):
        ring = [point(index, level / maximum) for index in range(count)]
        shape.draw_polyline(ring + ring[:1])
        shape.finish(color=grid, width=0.9 if level == maximum else 0.5,
                     fill=(0.95, 0.975, 0.97) if level == maximum else None)
    for index in range(count):
        shape.draw_line(center, point(index, 1))
        shape.finish(color=grid, width=0.5)
    spots = {index: point(index, min(max(values[index], 0), maximum) / maximum) for index in known}
    if len(known) == count:
        shape.draw_polyline([spots[index] for index in range(count)] + [spots[0]])
        shape.finish(color=brand, width=1.6, fill=brand, fill_opacity=0.16)
    else:
        for index in known:
            if (index + 1) % count in spots:
                shape.draw_line(spots[index], spots[(index + 1) % count])
                shape.finish(color=brand, width=1.6)
    for spot in spots.values():
        shape.draw_circle(spot, 2.4)
        shape.finish(color=brand, fill=brand, width=0.5)
    shape.commit()
    for index, label in enumerate(labels):
        shown = "证据不足" if values[index] is None else f"{values[index]:g} 分"
        across, down = math.cos(angle(index)), math.sin(angle(index))
        anchor = point(index, 1)
        y = anchor.y + (-20 if down < -0.3 else 16 if down > 0.3 else -2)
        # 维度名称和得分写在顶点外侧：上下居中，左右两侧分别向外对齐。
        for line, size, color, drop in ((label, 10, ink, 0), (shown, 9, muted if values[index] is None else brand, 12)):
            parts = _font_runs(line)
            width = sum(fitz.get_text_length(text, font, size) for font, text in parts)
            x = anchor.x - width / 2 if abs(across) < 0.3 else anchor.x + 9 if across > 0 else anchor.x - 9 - width
            for font, text in parts:
                page.insert_text(fitz.Point(x, y + drop), text, fontsize=size, fontname=font, color=color)
                x += fitz.get_text_length(text, font, size)
    return page.get_pixmap(dpi=200).tobytes("png")


def _font_runs(text):
    """把文字按中文、西文分段：内置中文字体里数字和字母的间距偏大，西文部分改用西文字体。"""
    parts = []
    for char in text:
        font = "helv" if char.isascii() else "china-s"
        if parts and parts[-1][0] == font:
            parts[-1][1] += char
        else:
            parts.append([font, char])
    return parts


def _radar(doc, evaluations):
    picture = radar_png([item.get("dimension", "") for item in evaluations],
                        [_number(item.get("score")) for item in evaluations])
    if not picture:
        return
    holder = style.paragraph(doc, "", after=0, line=1.0, align="center", keep=True)
    holder.add_run().add_picture(io.BytesIO(picture), width=Twips(7030))
    style.paragraph(doc, f"六维辅助评分雷达图（满分 {FULL_SCORE} 分，由内向外每格 1 分）", size=9, color=style.MUTED,
                    align="center", after=8)


# ---- 面试报告、综合报告 ------------------------------------------------------------

def _summary_card(doc, metadata, evaluations, match, content, comprehensive):
    """结果摘要卡：辅助匹配度、六维评分与第二阶段材料的数量、报告版本。

    报告只在人工确认之后才能导出，所以不再写报告状态。
    """
    facts = []
    if evaluations:
        scored = sum(item.get("score") is not None for item in evaluations)
        facts.append(("六维评分", f"已形成评分 {scored}/{len(evaluations)} 项"))
    if comprehensive:
        standalone = content.get("mode") == "standalone"
        facts.append(("人工面试转写" if standalone else "第二阶段转写", f"{len(content.get('stage_two_evidence') or [])} 段"))
        facts.append(("综合分析" if standalone else "两阶段综合分析", f"{len(content.get('findings') or [])} 项"))
    facts.append(("报告版本", f"v{metadata.get('version', '—')}"))
    if match.get("score") is not None:
        score = match["score"]
        metric = style.metric("第一阶段辅助匹配度" if comprehensive else "辅助匹配度",
                              f"{score:g}" if _number(score) is not None else score, "%")
        if len(facts) == 2:
            # 面试报告只有六维评分和版本两项，和匹配度放在同一行。
            rows = [[(5, metric), *((3, style.fact(label, value)) for label, value in facts)]]
        else:
            rows = [[(5, metric), (6, style.inline_fact(*facts.pop(0)))], [(1, style.fact(label, value)) for label, value in facts]]
    else:
        rows = [[(1, style.fact(label, value)) for label, value in facts]]
    style.card(doc, rows)


def _evaluation_table(doc, evaluations):
    widths = style.spread((4.6, 1.8, 5.6, 2.2, 2.8), style.content_width(doc))
    lines = [style.header_cells(("评价维度", "权重", "辅助评分", "评价等级", "证据充分程度"), widths, (1, 2, 3, 4))]
    for item in evaluations:
        score = item.get("score")
        if score is None:
            bar = style.bar_cells(widths[2], None, "证据不足")
        elif _number(score) is None:
            bar = [style.cell(str(score), widths[2], align="center")]
        else:
            bar = style.bar_cells(widths[2], score, f"{score:g} / {FULL_SCORE}", maximum=FULL_SCORE)
        lines.append([
            style.cell(item.get("dimension", ""), widths[0]),
            style.cell(f"{item.get('weight', '—')}%", widths[1], align="center"), *bar,
            style.cell(item.get("level", "证据不足"), widths[3], align="center"),
            style.cell(item.get("evidence_sufficiency", "不足"), widths[4], align="center"),
        ])
    style.grid_table(doc, lines, header_rows=1, together=len(lines) <= 12)


def _evaluation_details(doc, evaluations):
    style.paragraph(doc, "以下逐项列出评分理由；左侧竖线标出的是回答原文摘录，其下注明关联回答和证据充分程度。",
                    size=9.5, color=style.MUTED)
    for index, item in enumerate(evaluations, start=1):
        style.heading(doc, f"{index}. {item.get('dimension', '')}", 3)
        style.paragraph(doc, [("评分理由：", 10.5, True, style.INK),
                              (str(item.get("reason") or "证据不足"), 10.5, False, style.INK)], line=1.3)
        answers = "、".join(str(value) for value in item.get("answer_ids", []) if value) or "证据不足"
        excerpt = str(item.get("evidence_excerpt") or "证据不足")
        sufficiency = item.get("evidence_sufficiency", "不足")
        if excerpt == "证据不足":
            # 没有可引用的回答原文时只写说明，不画引用线。
            style.paragraph(doc, f"证据摘录：证据不足　　关联回答：{answers}　　证据充分程度：{sufficiency}",
                            size=9.5, color=style.MUTED)
        else:
            style.quote(doc, excerpt, f"关联回答 {answers} · 证据充分程度：{sufficiency}")


def _match_table(doc, match):
    width = style.content_width(doc)
    score = match.get("score", "—")
    if _number(score) is None:
        value = [style.cell(f"{score}%", width - 1900)]
    else:
        value = style.bar_cells(width - 1900, score, f"{score:g}%")
    lines = [[style.cell("辅助匹配度", 1900, bold=True, fill=style.LIGHT), *value]]
    for label, key in (("匹配点", "matched"), ("待核实事项", "to_verify")):
        lines.append([style.cell(label, 1900, bold=True, fill=style.LIGHT),
                      style.cell("\n".join(str(entry) for entry in match.get(key, []) or ["暂无"]), width - 1900)])
    style.grid_table(doc, lines)


def _questions_and_answers(doc, records):
    for index, item in enumerate(records, start=1):
        code = item.get("answer_code") or f"A{index}"
        style.heading(doc, f"{code}  {item.get('question', '')}", 3)
        _transcript(doc, "主体回答：", item.get("transcript") or "未作答")
        followup = item.get("followup") or item.get("followup_text")
        if followup:
            style.paragraph(doc, f"动态追问：{followup}", size=10, color=style.MUTED, after=3, line=1.3)
            _transcript(doc, "追问回答：", item.get("followup_answer") or "未作答")


def _capture_label(value):
    """录音编号用于回查录音；完整编号很长，报告里只写前 8 位。"""
    value = str(value or "")
    return (value[:8] + "…") if len(value) > 16 else (value or "—")


def _spoken_rows(evidence):
    """第二阶段转写的表格行。已经区分发言人的连续对话逐段列出，并标明是推断还是面试官确认过。"""
    rows = []
    for index, item in enumerate(evidence, start=1):
        turns = item.get("turns") if item.get("speaker") == "dialogue" else None
        if not turns:
            rows.append([index, SPEAKERS.get(item.get("speaker"), "发言人待核对"), item.get("text") or "暂无有效转写",
                         _capture_label(item.get("id"))])
            continue
        for order, turn in enumerate(turns, start=1):
            mark = "已确认" if turn.get("basis") == "confirmed" else "推断"
            rows.append([f"{index}-{order}", f"{ROLES.get(turn.get('role'), ROLES['unknown'])}\n（{mark}）",
                         turn.get("text") or "", _capture_label(item.get("id"))])
    return rows


def report_docx(title, content, metadata=None):
    metadata = metadata or {}
    candidate_name = style.person_label(metadata.get("candidate_name") or content.get("candidate") or "—")
    stage = content.get("stage_one") if isinstance(content.get("stage_one"), dict) else {}
    stage_content = stage.get("content", {}) if isinstance(stage, dict) else {}
    is_comprehensive = bool(stage_content or content.get("stage_two_evidence") or content.get("findings"))
    # 仅人工面试：候选人没有做第一阶段AI面试，综合报告依据简历、岗位要求和人工面试的旁听记录。
    standalone = is_comprehensive and content.get("mode") == "standalone"
    analysis_name = "综合分析" if standalone else "两阶段综合分析"
    evaluations = content.get("stage_one_evaluations") or stage_content.get("evaluations") or content.get("evaluations") or []
    match = content.get("stage_one_match_analysis") or stage_content.get("match_analysis") or content.get("match_analysis") or {}

    doc = style.new_document(title, title)
    style.info_line(doc, [("候选人", candidate_name), ("候选人编号", metadata.get("candidate_code", "—")),
                          ("应聘岗位", metadata.get("job_name", "—")), ("招聘批次", metadata.get("batch_name", "—")),
                          ("导出时间", style.now_text())])
    _summary_card(doc, metadata, evaluations, match, content, is_comprehensive)
    style.note(doc, NOTICE, label="说明：")
    if content.get("generation_mode") == PRESERVED:
        reason = content.get("generation_note")
        style.note(doc, "本报告生成时没有得到模型的分析结果" + (f"（{reason}）" if reason else "")
                   + f"，“{analysis_name}”中只保全了原文证据，结论由人工核实后补充。", label="注意：")

    number = 0

    def add_section(name):
        nonlocal number
        number += 1
        _section(doc, number, name)

    add_section("分析依据" if standalone else "第一阶段AI面试总结" if is_comprehensive else "面试纪要")
    style.paragraph(doc, content.get("stage_one_summary") or stage_content.get("summary") or content.get("summary") or "暂无纪要",
                    line=1.35)

    if evaluations:
        add_section("第一阶段六维评价" if is_comprehensive else "六维评价")
        _evaluation_table(doc, evaluations)
        _radar(doc, evaluations)
        _evaluation_details(doc, evaluations)

    stage_two_summary = content.get("stage_two_summary")
    stage_two_evidence = content.get("stage_two_evidence") or []
    if stage_two_summary or stage_two_evidence:
        add_section("人工面试旁听纪要" if standalone else "第二阶段AI旁听纪要")
        style.paragraph(doc, stage_two_summary or "已完成旁听，具体转写见下表。", line=1.35)
        if stage_two_evidence:
            if any(item.get("turns") for item in stage_two_evidence):
                style.paragraph(doc, "连续对话由系统按对话语境逐段区分发言人：标“推断”的是系统的判断，须核对录音；"
                                "标“已确认”的由面试官确认过。序号“1-2”表示第 1 段录音中的第 2 段发言。",
                                size=9.5, color=style.MUTED)
            style.table(doc, ["序号", "发言角色", "原始转写", "录音编号"], _spoken_rows(stage_two_evidence),
                        [1.4, 2.6, 10.2, 2.8], (0, 1, 3), split=True)

    findings = content.get("findings") or []
    if findings:
        add_section(analysis_name)
        style.table(doc, ["维度", "状态", "综合分析", "候选人回答原文" if standalone else "第二阶段证据"], [
            [f.get("dimension", ""), f.get("status", ""),
             f.get("analysis", "") + ("\n简历依据：" + f["resume_excerpt"] if f.get("resume_excerpt") else ""),
             (f.get("answer_excerpt") or "证据不足")
             + ("\n回答归属：" + f["speaker_basis"] if f.get("speaker_basis") else "")] for f in findings],
            [2.6, 3.0, 6.4, 5.0], (1,), split=True)

    resume = content.get("resume_context") or stage_content.get("resume_context") or {}
    checks = content.get("resume_checks") or stage_content.get("resume_checks") or []
    if resume or checks:
        add_section("简历与回答核对")
        style.paragraph(doc, "简历用于核实经历，不直接作为评分证据；未提及不代表不属实。"
                        + ("简历较长，本次仅分析前24000字。" if resume.get("truncated") else ""), line=1.35)
        if checks:
            style.table(doc, ["简历原文", "回答原文", "核对结果", "待核实说明"], [
                [c.get("resume_excerpt", ""), c.get("answer_excerpt") or "尚未覆盖", c.get("status", "待核实"), c.get("finding", "")]
                for c in checks], [4.6, 4.6, 2.4, 5.4], (2,), split=True)
        else:
            style.paragraph(doc, "暂无简历对照项，请人工核实。" if resume.get("status") == "ready"
                            else "本次没有可用简历文字，报告依据面试回答分析。")

    if match:
        add_section("第一阶段岗位匹配分析" if is_comprehensive else "岗位匹配分析")
        _match_table(doc, match)

    risks = []
    for item in [*(content.get("stage_one_risks") or stage_content.get("risks") or []), *(content.get("risks") or [])]:
        if item not in risks:
            risks.append(item)
    if risks:
        add_section("综合风险提示" if is_comprehensive else "风险提示")
        style.numbered(doc, risks)

    offline = content.get("adopted_questions") or stage_content.get("offline_questions") or content.get("offline_questions") or []
    if offline:
        add_section("本场采用的提问" if standalone and content.get("adopted_questions")
                    else "第二阶段已采用问题" if is_comprehensive and content.get("adopted_questions") else "线下人工面试建议问题")
        style.numbered(doc, offline)

    records = content.get("stage_one_qa_records") or stage_content.get("qa_records") or content.get("qa_records") or []
    if records:
        add_section("第一阶段完整问答记录" if is_comprehensive else "完整问答记录")
        _questions_and_answers(doc, records)
    return style.save_bytes(doc)


# ---- 候选人对比 ------------------------------------------------------------------

def _person(item, stacked=False) -> str:
    """候选人姓名和编号；stacked 为真时分两行写，用于较窄的列。"""
    person = item.get("candidate", {})
    name, code = style.person_label(person.get("name", "")), person.get("code", "")
    return f"{name}\n{code}" if stacked else f"{name}（{code}）"


def _comparison_card(doc, items):
    scored = [item for item in items if item.get("scores")]
    ranked = [item for item in items if _number(item.get("match_score")) is not None]
    top = max(ranked, key=lambda item: item["match_score"], default=None)
    listened = sum(bool(item.get("stage_two")) for item in items)
    # 横向页面较矮，摘要卡只占一行。
    style.card(doc, [[
        (3, style.metric("对比人数", len(items), " 人")), (4, style.metric("已有六维评分", len(scored), " 人")),
        (5, style.fact("最高岗位匹配度", f"{_percent(top['match_score'])}　{_person(top)}" if top else "暂无")),
        (4, style.fact("人工面试旁听", f"已做 {listened} 人；未做 {len(items) - listened} 人")),
    ]])


def _score_table(doc, items, dimensions):
    count = len(dimensions)
    widths = style.spread([3.4, *([3.1] * count), 4.1], style.content_width(doc))
    lines = [style.header_cells(["候选人", *dimensions, "岗位匹配度"], widths, range(1, count + 2), size=9)]
    for item in items:
        scores = item.get("scores", {})
        line = [style.cell(_person(item, stacked=True), widths[0], size=9)]
        for name, width in zip(dimensions, widths[1:]):
            value = scores.get(name)
            line.append(style.cell(str(value) if value is not None else "证据不足", width, size=9, align="center",
                                   color=style.INK if value is not None else style.MUTED,
                                   fill=style.heat_fill(_number(value), FULL_SCORE)))
        line += style.bar_cells(widths[-1], _number(item.get("match_score")), _percent(item.get("match_score")), size=9)
        lines.append(line)
    style.grid_table(doc, lines, header_rows=1)
    style.paragraph(doc, f"六维评分满分 {FULL_SCORE} 分，底色越深表示得分越高；“证据不足”表示该维度没有形成评分，不按 0 分处理。",
                    size=9, color=style.MUTED)


def comparison_docx(job_name, items):
    """Export a readable comparison report with score and risk tables separated."""
    doc = style.new_document("候选人对比报告", "候选人对比报告", landscape=True)
    style.info_line(doc, [("岗位", job_name), ("对比人数", len(items)), ("导出时间", style.now_text())])
    _comparison_card(doc, items)
    style.paragraph(doc, f"本次对比共包含 {len(items)} 名候选人，数据来自各人的面试报告或综合报告。")
    style.note(doc, "对比结果仅供人工初筛参考，不构成自动录用或淘汰决定。", label="说明：")

    dimensions = []
    for item in items:
        for name in item.get("scores", {}):
            if name not in dimensions:
                dimensions.append(name)

    _section(doc, 1, "六维辅助评分对比")
    if dimensions:
        _score_table(doc, items, dimensions)
    else:
        style.paragraph(doc, "所选候选人没有做第一阶段AI面试，没有六维评分。", color=style.MUTED)

    _section(doc, 2, "岗位匹配度与人工面试旁听")
    style.table(doc, ["候选人", "岗位匹配度", "人工面试旁听"], [[
        _person(item), _percent(item.get("match_score")), "已做，有综合报告" if item.get("stage_two") else "未做",
    ] for item in items], [4.0, 2.4, 3.5], (1, 2))

    _section(doc, 3, "风险提示对比")
    style.table(doc, ["候选人", "风险提示"], [
        [_person(item), "\n".join(str(risk) for risk in item.get("risks", []) or ["暂无"])] for item in items],
        [3.0, 6.9], split=True)

    _section(doc, 4, "硬指标（简历档案核对，不计入评分）")
    records = []
    for item in items:
        hard = item.get("hard_indicators") or {}
        facts = "\n".join(f"{fact.get('label')}：{fact.get('value')}" for fact in hard.get("facts", []) if fact.get("value"))
        checks = "\n".join(f"{check.get('condition')}：{check.get('status')}" for check in hard.get("hard_checks", []))
        records.append([_person(item), facts or hard.get("reason") or "暂无档案", checks or "岗位未设置硬性条件"])
    style.table(doc, ["候选人", "档案事实", "岗位硬性条件核验"], records, [2.6, 4.3, 3.0], split=True)
    return style.save_bytes(doc)
