"""Rebuild disposable test documents from explicitly fictional UTF-8 records.

No original PDF/DOCX, document properties, logo, or customer attachment is shipped.
"""
from pathlib import Path
import json
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'backend/tests/matching/fixtures'


def build():
    destination = FIXTURES / 'demo_data'
    for record in json.loads((FIXTURES / 'synthetic_resumes.json').read_text()):
        path = destination / record['path']
        path.parent.mkdir(parents=True, exist_ok=True)
        document = pymupdf.open()
        page = document.new_page()
        remaining = page.insert_textbox(pymupdf.Rect(40, 40, 555, 800), record['text'],
                                        fontname='china-s', fontsize=10, lineheight=1.4)
        if remaining < 0:
            raise ValueError('Synthetic fixture exceeds one page')
        document.set_metadata({'title': 'Synthetic test fixture', 'author': 'TalentFlow demo'})
        document.save(path)
        document.close()
    jobs = json.loads((destination / 'jobs.json').read_text())
    for job in jobs:
        path = destination / 'job_descriptions' / (job['id'] + '_demo.docx')
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading(job['name'], 0)
        document.add_paragraph('仅用于系统演示，内容纯属虚构')
        for heading, field in [('岗位职责', 'responsibilities'), ('任职要求', 'requirements')]:
            document.add_heading(heading, 1)
            for item in job[field]: document.add_paragraph(item)
        document.core_properties.author = 'TalentFlow demo'
        document.save(path)


if __name__ == '__main__':
    build()
